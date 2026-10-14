"""Browserexport mit realer Tabellenkalkulation gegen Web-API und Referenz prüfen."""
import argparse,json,math,subprocess,sys,urllib.request
from pathlib import Path
from tempfile import TemporaryDirectory
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT));sys.path.insert(0,str(ROOT/'tests'))
from web.backend.app import parse_project
from engine_check import check_file

def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--url',default='http://127.0.0.1:8000')
    parser.add_argument('--soffice',default='soffice')
    parser.add_argument('--project',default='web/test-results/Grauwasser_Projekt.json')
    parser.add_argument('--workbook',default='web/test-results/Grauwasser_Bewertungstool.xlsx')
    args=parser.parse_args()
    payload=json.loads((ROOT/args.project).read_text());payload.pop('format',None);payload.pop('version',None)
    request=urllib.request.Request(args.url+'/api/evaluate',data=json.dumps(payload).encode(),headers={'Content-Type':'application/json'})
    api=json.load(urllib.request.urlopen(request,timeout=30));assert api['valid'],api.get('errors')
    p,errors=parse_project(payload['project']);assert not errors,errors
    with TemporaryDirectory(prefix='grauwasser-web-export-check-') as directory:
        directory=Path(directory)
        cmd=[args.soffice,'-env:UserInstallation='+(directory/'profile').as_uri(),'--headless','--convert-to','xlsx','--outdir',str(directory),str((ROOT/args.workbook).resolve())]
        converted=subprocess.run(cmd,text=True,capture_output=True,timeout=180)
        assert converted.returncode==0,converted.stderr
        path=directory/Path(args.workbook).name;assert path.exists(),converted.stdout+converted.stderr
        result=check_file(path,p,payload['calendar'],False)
        w=load_workbook(path,data_only=True)
        for row,case in list(zip((8,9,10),api['scenarios']))+list(zip(range(14,26),api['sensitivity'])):
            if case['result']['valid']:
                for column,key in [(9,'used'),(10,'quota'),(11,'cashflow'),(12,'npv')]:
                    assert math.isclose(w['Szenarien'].cell(row,column).value,case['result'][key],rel_tol=1e-9,abs_tol=1e-7),(row,key)
                assert w['Szenarien'].cell(row,16).value==case['result']['overall']
            else: assert w['Szenarien'].cell(row,16).value=='Nicht bewertbar'
        for i,rating in enumerate(api['study']['ratings'],6):
            assert w['Testauswertung'].cell(i,2).value==rating['n']
            actual=w['Testauswertung'].cell(i,3).value
            if rating['mean'] is None: assert actual is None or actual==''
            else: assert math.isclose(actual,rating['mean'],abs_tol=1e-9)
    report={'status':'passed','engine':subprocess.check_output([args.soffice,'--version'],text=True).strip(),
            'description':'Browserexport gegen Web-API, unabhängige Referenz, 15 Varianten und Studienmittelwerte geprüft',**result}
    (ROOT/'web/test-results/export-engine-results.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(report,ensure_ascii=False))

if __name__=='__main__':main()
