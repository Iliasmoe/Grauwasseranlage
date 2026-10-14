"""Zwischenstände, Browserneustart, Entwürfe und fehlender localStorage, ohne Netz."""
import json, sys
from pathlib import Path
from tempfile import TemporaryDirectory
from playwright.sync_api import sync_playwright, expect
ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT))
URL='http://localhost:8766/Grauwasser_App.html'
records=[];errors=[];requests=[]
def passed(name):records.append({'case':name,'status':'passed'});print('PASS',name,flush=True)
def serve(ctx):
    content=(ROOT/'Grauwasser_App.html').read_bytes()
    def route(r):
        if r.request.url.split('#')[0]==URL:r.fulfill(body=content,content_type='text/html')
        else:requests.append(r.request.url);r.abort()
    ctx.route('**/*',route);ctx.set_offline(True)
def ready(p):expect(p.get_by_text('Tagesbilanz vollständig berechnet',exact=True)).to_be_visible(timeout=15000)
def inputs(p):p.get_by_role('button',name='Projekt & Anlage',exact=True).click();p.get_by_role('button',name='Wirtschaftlichkeit',exact=True).click();return p.get_by_label('Vermiedener Trinkwasserpreis (€/m³)',exact=True)
with sync_playwright() as pw, TemporaryDirectory(prefix='grauwasser-browser-profile-') as profile:
    def launch():
        ctx=pw.chromium.launch_persistent_context(profile,executable_path='/usr/bin/chromium',headless=True,args=['--no-sandbox'],viewport={'width':1440,'height':1080},accept_downloads=True);serve(ctx)
        p=ctx.pages[0];p.on('pageerror',lambda e:errors.append(str(e)));p.goto(URL,wait_until='networkidle');ready(p);return ctx,p
    ctx,p=launch();water=inputs(p);water.fill('4.25');ready(p)
    p.get_by_role('button',name='Zwischenstände',exact=True).click();d=p.get_by_role('dialog',name='Zwischenstände',exact=True)
    expect(d.locator('.checkpoint-row')).to_have_count(1)
    d.get_by_label('Bezeichnung des Zwischenstands').fill('Variante vor Preisänderung');d.get_by_role('button',name='Zwischenstand sichern',exact=True).click()
    expect(d.locator('.checkpoint-row')).to_have_count(2);d.get_by_role('button',name='Schließen',exact=True).click()
    water.fill('9.75');ready(p)
    p.get_by_role('button',name='Zwischenstände',exact=True).click();d=p.get_by_role('dialog',name='Zwischenstände',exact=True)
    row=d.locator('.checkpoint-row').filter(has_text='Variante vor Preisänderung');row.get_by_role('button',name='Wiederherstellen',exact=True).click();ready(p);expect(water).to_have_value('4.25')
    passed('Automatische und benannte Zwischenstände; Wiederherstellung korrekt')
    p.get_by_role('button',name='Zwischenstände',exact=True).click();d=p.get_by_role('dialog',name='Zwischenstände',exact=True)
    d.locator('.checkpoint-row').filter(has_text='Vor Wiederherstellung').get_by_role('button',name='Wiederherstellen',exact=True).click();ready(p);expect(water).to_have_value('9.75')
    passed('Stand vor Wiederherstellung bleibt ebenfalls verfügbar')
    water.fill('');expect(p.get_by_role('alert')).to_contain_text('Trinkwasserpreis');p.reload(wait_until='networkidle');water=inputs(p);expect(water).to_have_value('');expect(p.get_by_role('alert')).to_contain_text('Trinkwasserpreis')
    expect(p.get_by_role('button',name='Excel exportieren',exact=True)).to_be_disabled();ctx.close()
    ctx=pw.chromium.launch_persistent_context(profile,executable_path='/usr/bin/chromium',headless=True,args=['--no-sandbox'],viewport={'width':1440,'height':1080});serve(ctx)
    p=ctx.pages[0];p.on('pageerror',lambda e:errors.append(str(e)));p.goto(URL,wait_until='networkidle');water=inputs(p);expect(water).to_have_value('');expect(p.get_by_role('alert')).to_contain_text('Trinkwasserpreis')
    p.get_by_role('button',name='Zwischenstände',exact=True).click();d=p.get_by_role('dialog',name='Zwischenstände',exact=True);expect(d.locator('.checkpoint-row').filter(has_text='Variante vor Preisänderung')).to_have_count(1);d.get_by_role('button',name='Schließen',exact=True).click()
    passed('Unvollständiger Entwurf und Historie überleben Neuladen und vollständigen Browserneustart')
    water.fill('4.25');ready(p)
    p.get_by_role('button',name='Zwischenstände',exact=True).click();d=p.get_by_role('dialog',name='Zwischenstände',exact=True)
    for i in range(24):
        d.get_by_label('Bezeichnung des Zwischenstands').fill('Teststand '+str(i));d.get_by_role('button',name='Zwischenstand sichern',exact=True).click()
    expect(d.locator('.checkpoint-row')).to_have_count(20);expect(d.locator('.checkpoint-row').first).to_contain_text('Teststand 23')
    d.get_by_role('button',name='Schließen',exact=True).click()
    passed('Historie ist auf die letzten 20 Zwischenstände pro Projekt begrenzt')
    p.get_by_role('button',name='Neues Projekt',exact=True).click();d=p.get_by_role('dialog');d.get_by_label('Projektname',exact=True).fill('Unabhängiges Projekt');d.get_by_role('button',name='Projekt anlegen',exact=True).click();ready(p)
    p.get_by_role('button',name='Zwischenstände',exact=True).click();d=p.get_by_role('dialog',name='Zwischenstände',exact=True);expect(d.locator('.checkpoint-row')).to_have_count(0);d.get_by_role('button',name='Schließen',exact=True).click()
    p.get_by_label('Aktives Projekt').select_option(index=0);ready(p)
    p.get_by_role('button',name='Zwischenstände',exact=True).click();d=p.get_by_role('dialog',name='Zwischenstände',exact=True);expect(d.locator('.checkpoint-row')).to_have_count(20);d.get_by_role('button',name='Schließen',exact=True).click()
    passed('Zwischenstände bleiben ihrem Projekt zugeordnet')
    p.get_by_role('button',name='Methodik & KPIs',exact=True).click()
    with p.expect_download() as download:p.get_by_role('link',name='Methodik-Dokumentation').click()
    with TemporaryDirectory() as directory:
        path=Path(directory)/'Methodik.md';download.value.save_as(path);assert path.read_text()==(ROOT/'Methodik_Dokumentation.md').read_text()
    passed('Eingebettete Methodik ist ohne Netzwerk herunterladbar')
    p.get_by_role('button',name='Projektbewertung',exact=True).click();ready(p);p.screenshot(path=str(ROOT/'web/validation-results/html-dashboard.png'),full_page=False)
    p.set_viewport_size({'width':390,'height':844});p.get_by_role('button',name='Zwischenstände',exact=True).click();d=p.get_by_role('dialog',name='Zwischenstände',exact=True)
    assert not p.evaluate('document.documentElement.scrollWidth>innerWidth');expect(d.get_by_role('button',name='Zwischenstand sichern',exact=True)).to_be_visible();d.get_by_role('button',name='Schließen',exact=True).click()
    passed('Zwischenstand-Dialog funktioniert bei 390 px ohne Seitenüberlauf')
    p.set_viewport_size({'width':1440,'height':1080})
    # Neue Browser-XLSX-Schreiblogik muss Nutzertexte als Texte erhalten.
    from web.backend.app import defaults
    from openpyxl import load_workbook
    fixture=defaults();fixture['project']['name']='=1+1';fixture['project']['sources'][0]['name']='<script>alert(1)</script>'
    fixture['scenarios'][0]['name']='@SUM(A1:A2)'
    fixture['studies']=[{'id':'+1+1','role':'=2+2','experience':1,'minutes':10,'mistakes':0,'ratings':[4]*7,'tasks':['Erfüllt']*4,'errors':'=HYPERLINK("https://example.com")','comment':'<b>Text</b>'}]
    fixture={'format':'grauwasser-web','version':1,**fixture}
    p.locator('input[type=file]').set_input_files({'name':'literal.json','mimeType':'application/json','buffer':json.dumps(fixture).encode()});ready(p)
    with p.expect_download() as download:p.get_by_role('button',name='Excel exportieren',exact=True).click()
    with TemporaryDirectory() as directory:
        path=Path(directory)/'literal.xlsx';download.value.save_as(path);wb=load_workbook(path)
        for sheet,cell,value in [('Eingaben','B5','=1+1'),('Nutzungsprofile','A6','<script>alert(1)</script>'),('Szenarien','A8','@SUM(A1:A2)'),('Nutzertest','A6','+1+1'),('Nutzertest','R6','<b>Text</b>')]:
            assert wb[sheet][cell].value==value and wb[sheet][cell].data_type=='s',(sheet,cell)
        assert wb['Nutzertest']['A7'].value is None
        cached=load_workbook(path,data_only=True)
        for sheet in wb:
            for row in sheet:
                for cell in row:
                    if cell.data_type=='f':assert cached[sheet.title][cell.coordinate].value is None,(sheet.title,cell.coordinate)
    passed('Excel-Export: Nutzertexte bleiben Literaltexte; keine alten Formel-Caches oder Studienzeilen')
    ctx.close()
    browser=pw.chromium.launch(executable_path='/usr/bin/chromium',headless=True,args=['--no-sandbox'])
    ctx=browser.new_context();serve(ctx);ctx.add_init_script("Object.defineProperty(window,'localStorage',{get(){throw new DOMException('Blocked','SecurityError')}})")
    p=ctx.new_page();p.on('pageerror',lambda e:errors.append(str(e)));p.goto(URL,wait_until='networkidle');ready(p)
    expect(p.get_by_text('Nicht gespeichert',exact=True)).to_be_visible();expect(p.get_by_text('Browser-Speicher nicht verfügbar. Bitte Projekt als JSON sichern.',exact=True)).to_be_visible()
    water=inputs(p);water.fill('6.5');ready(p)
    with p.expect_download() as download:p.get_by_role('button',name='Projektdatei',exact=True).click()
    with TemporaryDirectory() as directory:
        path=Path(directory)/'backup.json';download.value.save_as(path);assert json.loads(path.read_text())['project']['water_price']==6.5
    passed('Blockierter Browser-Speicher: ehrlicher Hinweis, Weiterarbeiten und JSON-Sicherung möglich')
    browser.close()
assert not errors,errors
assert not requests,requests
report={'cases':records,'browser_errors':errors,'network_requests_beyond_document':requests,'delivery':'Unveränderte HTML-Datei an lokaler Testadresse; alle weiteren Anfragen gesperrt; offline=True','file_url_test':'Cloud-Browser blockiert file:// administrativ; keine Richtlinien geändert'}
(ROOT/'web/validation-results/offline-storage-results.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
print(str(len(records))+' Abläufe bestanden.')
