"""Stateless API: exakt dieselbe geprüfte Rechenbasis wie das Excel-Werkzeug."""
from __future__ import annotations
from dataclasses import asdict, fields
from datetime import date, timedelta
from io import BytesIO
from math import isfinite
from pathlib import Path
from statistics import mean, stdev
from tempfile import TemporaryDirectory
from typing import Any
import sys

from fastapi import FastAPI, HTTPException, Request
from fastapi.responses import FileResponse, JSONResponse, Response
from fastapi.exceptions import RequestValidationError
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, ConfigDict, Field, StrictInt, StrictFloat, StrictStr, StrictBool

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'src'))
from model import Project, Profile, INPUT_ROWS, default_calendar, evaluate

app = FastAPI(title='Grauwasser · Webanwendung', version='1.0.0', docs_url=None, redoc_url=None, openapi_url='/api/openapi.json')

SCENARIOS = [
    {'name': 'Konservativ', 'factors': [1.2, .9, .9, .8, 1.2, .01]},
    {'name': 'Basis', 'factors': [1, 1, 1, 1, 1, 0]},
    {'name': 'Optimistisch', 'factors': [.8, 1.1, 1.1, 1.2, .8, -.01]},
]
LABELS = {
    'name': 'Projektname', 'building': 'Gebäudetyp', 'location': 'Standort', 'plant': 'Anlage',
    'users': 'Nutzerzahl', 'year': 'Bilanzjahr', 'horizon': 'Betrachtungszeitraum',
    'other_water': 'Sonstiger Trinkwasserbedarf', 'yield_rate': 'Aufbereitungsausbeute',
    'availability': 'Verfügbarkeit', 'capacity': 'Aufbereitungskapazität', 'storage': 'Speichervolumen',
    'overlap': 'Direkt nutzbarer Tagesanteil', 'storage_loss': 'Übertragsverlust',
    'hygiene': 'Hygienische Nachweise', 'technical': 'Technische Nachweise', 'legal': 'Rechtliche Nachweise',
    'evidence': 'Nachweisbelege', 'planning': 'Planungskosten', 'equipment': 'Anlagenkosten',
    'installation': 'Montagekosten', 'extra_investment': 'Sonstige Investition', 'grant': 'Fördermittel',
    'maintenance': 'Wartungskosten', 'energy_intensity': 'Aufbereitungsenergie', 'energy_price': 'Strompreis',
    'other_cost': 'Sonstige Betriebskosten', 'water_price': 'Trinkwasserpreis', 'sewage_price': 'Abwassergebühr',
    'sewage_credit': 'Abwasser-Einsparanteil', 'discount': 'Kalkulationszins', 'water_growth': 'Wasserpreissteigerung',
    'cost_growth': 'Betriebskostensteigerung', 'energy_growth': 'Strompreissteigerung',
    'replacement_year': 'Ersatzjahr', 'replacement_cost': 'Ersatzinvestition', 'residual': 'Restwert',
    'sewage_growth': 'Abwasserpreissteigerung', 'billing_evidence': 'Gebührenbeleg', 'data_status': 'Datenstatus',
    'resource_green': 'Ressourcen: grüne Schwelle', 'resource_yellow': 'Ressourcen: gelbe Schwelle',
    'economic_green': 'Ökonomie: grüne Schwelle', 'economic_yellow': 'Ökonomie: gelbe Schwelle',
}

class Scenario(BaseModel):
    model_config = ConfigDict(extra='forbid', allow_inf_nan=False)
    name: str = Field(min_length=1, max_length=80)
    factors: list[float | None] = Field(min_length=6, max_length=6)

class StudyRecord(BaseModel):
    model_config = ConfigDict(extra='forbid', allow_inf_nan=False)
    id: StrictStr = Field(default='', max_length=80)
    role: StrictStr = Field(default='', max_length=200)
    experience: StrictInt | StrictFloat | StrictStr | StrictBool | None = None
    minutes: StrictInt | StrictFloat | StrictStr | StrictBool | None = None
    mistakes: StrictInt | StrictFloat | StrictStr | StrictBool | None = None
    ratings: list[StrictInt | StrictFloat | StrictStr | StrictBool | None] = Field(default_factory=list, max_length=7)
    tasks: list[StrictStr] = Field(default_factory=list, max_length=4)
    errors: StrictStr = Field(default='', max_length=4000)
    comment: StrictStr = Field(default='', max_length=4000)

class EvaluationInput(BaseModel):
    model_config = ConfigDict(extra='forbid')
    project: dict[str, Any]
    calendar: list[list[Any]] | None = Field(default=None, max_length=366)
    calendar_year: int | None = None
    scenarios: list[Scenario] = Field(default_factory=lambda: [Scenario(**x) for x in SCENARIOS], max_length=3, min_length=3)
    studies: list[StudyRecord] = Field(default_factory=list, max_length=50)


def problem(path: str, message: str) -> dict[str, str]:
    return {'path': path, 'message': message}


def parse_project(raw: dict[str, Any]) -> tuple[Project, list[dict[str, str]]]:
    p = Project()
    expected = {f.name for f in fields(Project)}
    errors = []
    for key in sorted(set(raw) - expected):
        errors.append(problem(key, 'Unbekannter Projektparameter.'))
    for key in sorted(expected):
        if key not in raw:
            errors.append(problem(key, f'{LABELS.get(key, key)} fehlt.'))
            continue
        value = raw[key]
        if key in ('sources', 'demands'):
            items = []
            if not isinstance(value, list) or len(value) > 6:
                errors.append(problem(key, 'Bitte höchstens sechs Nutzungsprofile angeben.'))
                continue
            for i, item in enumerate(value):
                prefix = f'{key}.{i}'
                profile_fields = {'name', 'group', 'users', 'frequency', 'liters', 'days'}
                if not isinstance(item, dict) or set(item) != profile_fields:
                    errors.append(problem(prefix, 'Ein Profil benötigt Name, Gruppe, Nutzer, Häufigkeit, Liter und Tage.'))
                    continue
                good = True
                for name in ('name', 'group'):
                    if not isinstance(item[name], str) or not item[name].strip() or len(item[name]) > 200:
                        errors.append(problem(prefix + '.' + name, 'Eine gültige Bezeichnung ist erforderlich.')); good = False
                for name in ('users', 'frequency', 'liters', 'days'):
                    v = item[name]
                    if type(v) not in (int, float) or not isfinite(v):
                        errors.append(problem(prefix + '.' + name, 'Bitte eine endliche Zahl eingeben.')); good = False
                if good:
                    items.append(Profile(**item))
            setattr(p, key, items)
        else:
            prototype = getattr(p, key)
            if isinstance(prototype, str):
                if not isinstance(value, str) or len(value) > 4000:
                    errors.append(problem(key, f'{LABELS[key]}: bitte Text eingeben.')); continue
            elif type(value) not in (int, float) or not isfinite(value):
                errors.append(problem(key, f'{LABELS[key]}: eine Zahl ist erforderlich.')); continue
            elif isinstance(prototype, int):
                if value != int(value):
                    errors.append(problem(key, f'{LABELS[key]} muss ganzzahlig sein.')); continue
                value = int(value)
            setattr(p, key, value)
    return p, errors


def model_errors(p: Project, names: list[str]) -> list[dict[str, str]]:
    messages = {
        'grant': 'Fördermittel dürfen die Gesamtinvestition nicht übersteigen.',
        'replacement_year': 'Ersatzjahr muss im Betrachtungszeitraum liegen; positive Ersatzkosten benötigen ein Jahr ab 1.',
        'billing_evidence': 'Positive Abwasserentlastung benötigt einen projektspezifischen Gebührenbeleg.',
        'resource thresholds': 'Ressourcenschwellen: 0 % ≤ Gelb ≤ Grün ≤ 100 %.',
        'economic thresholds': 'Ökonomieschwellen: −100 % ≤ Gelb ≤ Grün ≤ 100 %.',
        'calendar': 'Kalender muss alle Tage des Bilanzjahrs und ausschließlich 0/1-Werte enthalten.',
        'profile': 'Profile prüfen: Nutzer ≥ 0, Häufigkeit 0–100, Liter 0–1000 und ganzzahlige Tage 0–366.',
        'profile days': 'Profiltage dürfen bei positiver Menge die aktiven Kalendertage der Gruppe nicht übersteigen.',
        'scenario': 'Szenariofaktoren müssen zwischen 0 und 10 liegen; resultierender Zins zwischen 0 und 100 %.',
    }
    return [problem(key, messages.get(key, f'{LABELS.get(key, key)}: Wertebereich oder Pflichtfeld prüfen.')) for key in dict.fromkeys(names)]


def study_statistics(records: list[dict[str, Any]]) -> dict[str, Any]:
    identified = [r for r in records if isinstance(r.get('id'), str) and r['id'].strip()]
    ids = [r['id'].strip() for r in identified]
    ratings = []
    for dimension in range(7):
        values = []; missing = 0; invalid = 0
        for r in identified:
            answers = r.get('ratings', [])
            v = answers[dimension] if isinstance(answers, list) and dimension < len(answers) else None
            if v is None or v == '': missing += 1
            elif type(v) in (int, float) and isfinite(v) and v in (1, 2, 3, 4, 5): values.append(v)
            else: invalid += 1
        ratings.append({'n': len(values), 'mean': mean(values) if values else None,
                        'std': stdev(values) if len(values) > 1 else None,
                        'distribution': [values.count(i) for i in range(1, 6)], 'missing': missing, 'invalid': invalid})
    def valid_value(v, max_value, integer=False):
        return type(v) in (int, float) and isfinite(v) and 0 <= v <= max_value and (not integer or v == int(v))
    times = [r['minutes'] for r in identified if valid_value(r.get('minutes'), 1440)]
    mistakes = [r['mistakes'] for r in identified if valid_value(r.get('mistakes'), 10000, True)]
    def has_data(value):
        return any(has_data(v) for v in value) if isinstance(value, list) else value is not None and value != ''
    orphaned = sum(1 for r in records if r not in identified and any(has_data(v) for k, v in r.items() if k != 'id'))
    task_distribution = []
    for i in range(4):
        task_distribution.append({s: sum(1 for r in identified if isinstance(r.get('tasks'), list) and len(r['tasks']) > i and r['tasks'][i] == s) for s in ('Erfüllt', 'Teilweise', 'Nicht erfüllt')})
    return {'count': len(identified), 'duplicates': sorted({i for i in ids if ids.count(i) > 1}),
            'orphaned': orphaned, 'ratings': ratings, 'time_n': len(times), 'mean_minutes': mean(times) if times else None,
            'mistakes_n': len(mistakes), 'mean_mistakes': mean(mistakes) if mistakes else None, 'tasks': task_distribution}


def sensitivity_cases():
    cases = []
    for label, index in [('Investition', 0), ('Trinkwasserpreis', 1), ('Betriebswasserbedarf', 2), ('Grauwasseranfall', 3), ('Betriebskosten', 4), ('Kalkulationszins', 5)]:
        for direction in (-1, 1):
            factors = [1, 1, 1, 1, 1, 0]
            factors[index] = direction * .01 if index == 5 else 1 + direction * .2
            suffix = ('−1 %-Punkt' if direction < 0 else '+1 %-Punkt') if index == 5 else ('−20 %' if direction < 0 else '+20 %')
            cases.append({'name': label + ' ' + suffix, 'factors': factors})
    return cases


def compute(body: EvaluationInput) -> dict[str, Any]:
    p, errors = parse_project(body.project)
    if errors:
        return {'valid': False, 'errors': errors, 'study': study_statistics([r.model_dump() for r in body.studies])}
    if body.calendar_year is not None and body.calendar_year != p.year:
        return {'valid': False, 'errors': [problem('calendar', 'Der Kalender gehört zu einem anderen Jahr. Kalender für das Bilanzjahr neu erstellen oder importieren.')], 'study': study_statistics([r.model_dump() for r in body.studies])}
    if not 2020 <= p.year <= 2100:
        return {'valid': False, 'errors': [problem('year', 'Bilanzjahr muss zwischen 2020 und 2100 liegen.')], 'study': study_statistics([r.model_dump() for r in body.studies])}
    calendar = body.calendar if body.calendar is not None else default_calendar(p.year)
    if any(len(r) != 3 or any(type(v) is not int or v not in (0, 1) for v in r) for r in calendar):
        return {'valid': False, 'errors': [problem('calendar', 'Kalenderwerte müssen ganzzahlig 0 oder 1 sein.')], 'study': study_statistics([r.model_dump() for r in body.studies])}
    result = evaluate(p, calendar)
    if not result['valid']:
        return {'valid': False, 'errors': model_errors(p, result['errors']), 'study': study_statistics([r.model_dump() for r in body.studies])}
    warnings = []
    if p.data_status == 'Beispieldaten': warnings.append({'kind': 'info', 'text': 'Beispieldaten: Mengen, Preise und Anlagenparameter durch Projektbelege ersetzen.', 'page': 'inputs'})
    if result['gate'] != 'Grün': warnings.append({'kind': 'warning', 'text': 'Hygienische, technische oder rechtliche Nachweise sind ungeklärt bzw. nicht erfüllt. Keine Freigabe.', 'page': 'inputs'})
    if result['static'] is not None and result['static'] > p.horizon: warnings.append({'kind': 'warning', 'text': 'Statische Amortisation liegt außerhalb des Betrachtungszeitraums.', 'page': 'inputs'})
    if result['cashflow'] <= 0: warnings.append({'kind': 'warning', 'text': 'Im ersten Jahr entsteht keine positive Nettoeinsparung.', 'page': 'inputs'})
    if result['dynamic'] is not None and not result['stable_recovery']: warnings.append({'kind': 'warning', 'text': 'Nach erster dynamischer Amortisation tritt später erneut Unterdeckung auf.', 'page': 'inputs'})
    months = [{'month': i, 'raw': 0., 'used': 0., 'demand': 0., 'makeup': 0.} for i in range(1, 13)]
    for i, row in enumerate(result['daily']):
        day = date(p.year, 1, 1) + timedelta(days=i)
        row['date'] = day.isoformat()
        for key in ('raw', 'used', 'demand', 'makeup'): months[day.month - 1][key] += row[key]
    cumulative = []; value = 0
    for i, v in enumerate(result['discounted']):
        value += v; cumulative.append({'year': i, 'cashflow': result['flows'][i], 'discounted': v, 'cumulative': value})
    def run_case(s):
        result = evaluate(p, calendar, tuple(s['factors'])) if all(type(v) in (int, float) and isfinite(v) for v in s['factors']) else {'valid': False, 'errors': ['scenario'], 'overall': 'Nicht bewertbar'}
        result.pop('daily', None); result.pop('flows', None); result.pop('discounted', None)
        return {**s, 'result': result}
    return {'valid': True, 'errors': [], 'result': result, 'monthly': months, 'payments': cumulative,
            'scenarios': [run_case(s.model_dump()) for s in body.scenarios],
            'sensitivity': [run_case(s) for s in sensitivity_cases()],
            'warnings': warnings, 'calendar_counts': [sum(row[i] for row in calendar) for i in range(3)],
            'study': study_statistics([r.model_dump() for r in body.studies])}


@app.exception_handler(RequestValidationError)
async def invalid_request(request: Request, exc: RequestValidationError):
    return JSONResponse({'detail': 'Projektstruktur oder Eingabetyp prüfen.', 'errors': [{'path': '.'.join(str(x) for x in e['loc']), 'message': e['msg']} for e in exc.errors()]}, status_code=422)

@app.middleware('http')
async def headers(request: Request, call_next):
    length = request.headers.get('content-length')
    if length:
        try: size = int(length)
        except ValueError: return JSONResponse({'detail': 'Ungültige Content-Length.'}, status_code=400)
        if size > 2_000_000: return JSONResponse({'detail': 'Datei zu groß (maximal 2 MB).'}, status_code=413)
    response = await call_next(request)
    response.headers['X-Content-Type-Options'] = 'nosniff'
    response.headers['Referrer-Policy'] = 'same-origin'
    response.headers['Content-Security-Policy'] = "default-src 'self'; script-src 'self'; style-src 'self' 'unsafe-inline'; img-src 'self' data:; font-src 'self'; connect-src 'self'; frame-ancestors 'none'; base-uri 'self'"
    if request.url.path.startswith('/api/'):
        response.headers['Cache-Control'] = 'no-store'
    return response

@app.get('/api/health')
def health():
    return {'status': 'ok', 'model': 'greywater-daily-v1'}

@app.get('/api/default')
def defaults():
    p = Project()
    return {'project': asdict(p), 'calendar': default_calendar(p.year), 'calendar_year': p.year, 'scenarios': SCENARIOS, 'studies': []}

@app.get('/api/calendar/{year}')
def calendar(year: int):
    if not 2020 <= year <= 2100: raise HTTPException(422, 'Bilanzjahr 2020–2100 erforderlich.')
    return {'calendar': default_calendar(year), 'calendar_year': year}

@app.post('/api/evaluate')
def calculate(body: EvaluationInput):
    return compute(body)

@app.post('/api/export/xlsx')
def export_excel(body: EvaluationInput):
    result = compute(body)
    if not result['valid']: raise HTTPException(422, result['errors'])
    from generate import build
    from openpyxl import load_workbook
    p, _ = parse_project(body.project)
    schedule = body.calendar if body.calendar is not None else default_calendar(p.year)
    with TemporaryDirectory(prefix='grauwasser-export-') as directory:
        path = Path(directory) / 'Grauwasser_Bewertungstool.xlsx'
        build(path)
        wb = load_workbook(path)
        for key, row in INPUT_ROWS.items(): wb['Eingaben'].cell(row, 2).value = getattr(p, key)
        for start, profiles in [(6, p.sources), (17, p.demands)]:
            for i in range(6):
                item = profiles[i] if i < len(profiles) else Profile('Unbenutzt ' + str(i + 1), 'Individuell', 0, 0, 0, 0)
                for j, v in enumerate([item.name, item.group, item.users, item.frequency, item.liters, item.days], 1): wb['Nutzungsprofile'].cell(start + i, j).value = v
        for i in range(366):
            for j, v in enumerate(schedule[i] if i < len(schedule) else [0, 0, 0], 3): wb['Kalender'].cell(i + 5, j).value = v
        for row, key in [(5, 'economic_green'), (6, 'economic_yellow'), (9, 'resource_green'), (10, 'resource_yellow')]: wb['Schwellenwerte'].cell(row, 2).value = getattr(p, key)
        for row, s in zip((8, 9, 10), body.scenarios):
            name_cell = wb['Szenarien'].cell(row, 1); name_cell.value = s.name; name_cell.data_type = 's'; name_cell.quotePrefix = True
            for j, v in enumerate(s.factors, 2): wb['Szenarien'].cell(row, j).value = v
        for row, record in enumerate(body.studies, 6):
            s = record.model_dump()
            scalars = [s.get('id'), s.get('role'), s.get('experience'), s.get('minutes'), s.get('mistakes')]
            ratings = list(s.get('ratings', []))[:7]; ratings += [None] * (7 - len(ratings))
            tasks = list(s.get('tasks', []))[:4]; tasks += [None] * (4 - len(tasks))
            for j, v in enumerate(scalars + ratings + tasks + [s.get('errors'), s.get('comment')], 1):
                cell = wb['Nutzertest'].cell(row, j); cell.value = v
                if isinstance(v, str): cell.data_type = 's'; cell.quotePrefix = True
        # Nutzereingaben als Text erhalten, niemals als Formeln exportieren.
        for name in ('Eingaben', 'Nutzungsprofile', 'Szenarien'):
            for row in wb[name]:
                for cell in row:
                    if not cell.protection.locked and isinstance(cell.value, str) and cell.value.startswith(('=', '+', '-', '@')):
                        cell.data_type = 's'; cell.quotePrefix = True
        buffer = BytesIO(); wb.save(buffer)
    return Response(buffer.getvalue(), media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition': 'attachment; filename="Grauwasser_Bewertungstool.xlsx"'})

DOCUMENTS = {'methodik': 'Methodik_Dokumentation.md', 'probandentest': 'Probandentest.md', 'validierung': 'Validierungsbericht.md', 'anleitung': 'README.md'}

@app.get('/api/documents/{key}')
def document(key: str):
    if key not in DOCUMENTS: raise HTTPException(404, 'Dokument nicht gefunden.')
    return FileResponse(ROOT / DOCUMENTS[key], media_type='text/markdown; charset=utf-8', filename=DOCUMENTS[key])

DIST = ROOT / 'web/frontend/dist'
if DIST.is_dir():
    app.mount('/assets', StaticFiles(directory=DIST / 'assets'), name='assets')

@app.get('/')
def index():
    path = DIST / 'index.html'
    if not path.exists(): raise HTTPException(503, 'Frontend fehlt. In web/frontend zuerst npm ci und npm run build ausführen.')
    return FileResponse(path, media_type='text/html')
