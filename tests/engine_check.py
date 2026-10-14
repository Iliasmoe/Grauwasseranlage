"""Mutationstests der echten XLSX über LibreOffice; nie Python statt Excel bewerten."""
import argparse, json, math, subprocess, sys, time
from copy import deepcopy
from pathlib import Path
from tempfile import TemporaryDirectory
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'src')); sys.path.insert(0,str(ROOT/'tests'))
from model import Project, Profile, INPUT_ROWS, default_calendar, evaluate
from test_model import simple

RESULT_ROWS={'raw':6,'processed':7,'available':8,'demand':9,'used':10,'makeup':11,
'gross':13,'net_investment':14,'costs':16,'avoided':17,'cashflow':18,'npv':20,
'quota':22,'cover':23,'utilization':24,'specific_cost':25,'economic':31,
'resource':32,'overall':33,'overflow':26,'storage_losses':38,'end_storage':39}

def apply_project(wb,p,calendar):
    for key,row in INPUT_ROWS.items(): wb['Eingaben'].cell(row,2,getattr(p,key)) if getattr(p,key) is not None else setattr(wb['Eingaben'].cell(row,2),'value',None)
    ws=wb['Nutzungsprofile']
    for start,profiles in [(6,p.sources),(17,p.demands)]:
        for i in range(6):
            item=profiles[i] if i<len(profiles) else Profile('Unbenutzt '+str(i),'Individuell',0,0,0,0)
            for j,v in enumerate([item.name,item.group,item.users,item.frequency,item.liters,item.days],1): ws.cell(start+i,j,v)
    for i in range(366):
        for j,v in enumerate(calendar[i] if i<len(calendar) else (0,0,0),3): wb['Kalender'].cell(5+i,j,v)
    for row,key in [(5,'economic_green'),(6,'economic_yellow'),(9,'resource_green'),(10,'resource_yellow')]: wb['Schwellenwerte'].cell(row,2,getattr(p,key))

def cases():
    out=[]
    def add(name,p=None,c=None):
        p=p or Project(); c=c if c is not None else default_calendar(p.year if isinstance(p.year,int) else 2026)
        out.append((name,p,c))
    add('01_Basis')
    p=Project(); [setattr(x,'users',0) for x in p.sources]; add('02_Kein_Grauwasser',p)
    p=Project(); [setattr(x,'users',0) for x in p.demands]; add('03_Kein_Bedarf',p)
    p,c=simple(); add('04_Anfall_kleiner_Bedarf',p,c)
    p,c=simple(need=.5); add('05_Anfall_groesser_Bedarf',p,c)
    p=Project(); p.equipment=1e7; add('06_Hohe_Investition',p)
    p=Project(); p.water_price=0; add('07_Negativer_Cashflow',p)
    p,c=simple(investment=1460); add('08_Amortisation_innerhalb',p,c)
    p,c=simple(investment=14600); add('09_Amortisation_ausserhalb',p,c)
    p=Project(); p.water_price=None; add('10_Fehlende_Eingabe',p)
    p=Project(); p.yield_rate=1.5; add('11_Ungueltige_Eingabe',p)
    p,c=simple(); p.hygiene='Ungeklärt'; add('12_Hygiene_ungeklaert',p,c)
    p=Project(); p.resource_green=.8; p.resource_yellow=.3; add('13_Geaenderte_Schwellen',p)
    p=Project(); p.sources[0].days=100; add('14_Geaenderte_Tage',p)
    p,c=simple(investment=1000); p.discount=.05; add('15_Zins_Referenz',p,c)
    p,c=simple(raw=1,need=1); p.sources[0].group='Schulbetrieb'; p.sources[0].days=1; p.demands[0].group='Halle'; p.demands[0].days=1
    p.storage=1; p.overlap=0; p.storage_loss=.1; add('16_Tagesversatz_Speicher',p,[(1,0,0),(0,1,0)]+[(0,0,0)]*363)
    p,c=simple(); p.hygiene='Nicht erfüllt'; add('17_Hygiene_nicht_erfuellt',p,c)
    p=Project(); p.resource_green=.01; p.resource_yellow=.8; add('18_Ungueltige_Schwellen',p)
    p=Project(); p.year=2028; add('19_Schaltjahr',p)
    p=Project(); p.sewage_credit=1; add('20_Abwasser_unbelegt',p)
    p=Project(); p.sewage_credit=1; p.billing_evidence='Bestätigter Gebührenbescheid 123, variable Entlastung'; add('21_Abwasser_belegt',p)
    p,c=simple(); p.grant=100; add('22_Vollfoerderung',p,c)
    p,c=simple(investment=730); p.replacement_year=3; p.replacement_cost=10000; add('23_Spaetere_Unterdeckung',p,c)
    p,c=simple(raw=1,need=0); p.sources[0].days=1; p.storage=1; p.overlap=0; add('24_Endbestand',p,[(0,0,0)]*364+[(0,0,1)])
    p=Project(); p.water_price='ungueltig'; add('25_Text_statt_Zahl',p)
    p=Project(); p.horizon=2.5; add('26_Nicht_ganzzahlig',p)
    return out

def compare(actual,expected,label):
    if isinstance(expected,(float,int)) and not isinstance(expected,bool):
        assert isinstance(actual,(float,int)) and math.isclose(actual,expected,rel_tol=1e-9,abs_tol=1e-7),(label,actual,expected)
    else: assert actual==expected,(label,actual,expected)

def check_file(path,p,calendar,baseline=False):
    wb=load_workbook(path,data_only=True); ref=evaluate(p,calendar); ws=wb['Berechnung']
    errs=[(s.title,c.coordinate,c.value) for s in wb for row in s for c in row if c.data_type=='e']
    assert not errs,errs[:10]
    if not ref['valid']:
        assert ws['B5'].value>0,ref
        for c in ['B31','B32','B33']: compare(ws[c].value,'Nicht bewertbar',c)
        compare(wb['Dashboard']['A8'].value,'Eingaben prüfen','Dashboard Fehler')
    else:
        compare(ws['B5'].value,0,'Pflichtfehler')
        for key,row in RESULT_ROWS.items():
            expected=ref[key]
            if expected is None: expected='Keine Einsparmenge'
            compare(ws.cell(row,2).value,expected,key)
        compare(ws['B35'].value,0,'Volumenbilanz')
        compare(ws['B36'].value,0,'Profilintegration')
        compare(wb['Dashboard']['G8'].value,ref['used'],'Dashboard Einsparung')
        compare(wb['Dashboard']['G16'].value,ref['npv'],'Dashboard NPV')
        compare(wb['Dashboard']['I20'].value,ref['overall'],'Dashboard Ampel')
        static=ref['static'] if ref['static'] is not None else 'Keine Nettoinvestition' if ref['net_investment']==0 else 'Nicht erreichbar'
        dynamic=ref['dynamic'] if ref['dynamic'] is not None else 'Keine Nettoinvestition' if ref['net_investment']==0 else 'Im Zeitraum nicht erreicht'
        compare(ws['B19'].value,static,'Statisch'); compare(ws['B21'].value,dynamic,'Dynamisch')
        for row in wb['Wasserbilanz'].iter_rows(min_row=5,max_row=370):
            used=row[11].value; demand=row[4].value
            assert used>=-1e-9 and used<=demand+1e-8,(row[0].coordinate,used,demand)
            compare(row[15].value,0,'Tagesbilanz')
    if baseline:
        scenarios=wb['Szenarien']
        for r in [8,9,10]+list(range(14,26)):
            factors=tuple(scenarios.cell(r,c).value for c in range(2,8))
            result=evaluate(p,calendar,factors)
            compare(scenarios.cell(r,8).value,'OK','Variante gültig')
            for c,key in [(9,'used'),(10,'quota'),(11,'cashflow'),(12,'npv'),(14,'economic'),(15,'resource'),(16,'overall')]:
                compare(scenarios.cell(r,c).value,result[key],f'Variante {r} {key}')
    return {'valid':ref['valid'],'used':ref.get('used'),'npv':ref.get('npv'),'overall':ref.get('overall'),'excel_error_cells':len(errs)}

def main():
    parser=argparse.ArgumentParser(); parser.add_argument('--workbook',default='build/raw/Grauwasser_Bewertungstool.xlsx'); parser.add_argument('--output',default='Validierung_Ergebnisse.json'); parser.add_argument('--soffice',default='soffice'); args=parser.parse_args()
    began=time.time(); tests=cases(); records=[]
    with TemporaryDirectory(prefix='grauwasser-engine-') as directory:
        directory=Path(directory); inputs=directory/'input'; outputs=directory/'output'; inputs.mkdir(); outputs.mkdir()
        for name,p,calendar in tests:
            wb=load_workbook(ROOT/args.workbook); apply_project(wb,p,calendar)
            if name=='01_Basis':
                # Absichtlich gültige, fehlende und ungültige Antworten: n=2, Mittel=3.
                s=wb['Nutzertest']
                for row,uid,rating,t,errors in [(6,'P001',2,10,1),(7,'P002',4,20,3),(8,'P003',9,30,2),(9,'P004',None,None,None)]:
                    s.cell(row,1,uid); s.cell(row,4,t); s.cell(row,5,errors)
                    if rating is not None: s.cell(row,6,rating)
            wb.save(inputs/(name+'.xlsx')); print('prepared',name,flush=True)
        cmd=[args.soffice,'-env:UserInstallation='+(directory/'profile').as_uri(),'--headless','--convert-to','xlsx','--outdir',str(outputs)]+[str(inputs/(name+'.xlsx')) for name,_,_ in tests]
        proc=subprocess.run(cmd,text=True,capture_output=True,timeout=1200)
        print(proc.stdout,flush=True)
        if proc.returncode: raise RuntimeError(proc.stderr)
        for name,p,calendar in tests:
            path=outputs/(name+'.xlsx'); assert path.exists(),(name,proc.stderr)
            try:
                result=check_file(path,p,calendar,baseline=name=='01_Basis')
                if name=='01_Basis':
                    w=load_workbook(path,data_only=True)['Testauswertung']
                    for cell,value in [('B6',2),('C6',3),('D6',math.sqrt(2)),('J6',1),('K6',1),('B15',4),('B17',3),('B18',20),('B19',2)]: compare(w[cell].value,value,'Nutzertest '+cell)
                records.append({'case':name,'status':'passed',**result}); print('PASS',name,flush=True)
            except Exception as error:
                records.append({'case':name,'status':'failed','error':str(error)}); print('FAIL',name,str(error),flush=True)
    report={'engine':subprocess.check_output([args.soffice,'--version'],text=True).strip(),'elapsed_seconds':round(time.time()-began,1),'cases':records,'passed':sum(x['status']=='passed' for x in records),'failed':sum(x['status']=='failed' for x in records)}
    (ROOT/args.output).write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({k:v for k,v in report.items() if k!='cases'},ensure_ascii=False),flush=True)
    if report['failed']: raise SystemExit(1)

if __name__=='__main__': main()
