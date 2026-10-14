"""Strukturprüfung jeder Formel inklusive vollständiger Abhängigkeitsgraph."""
from bisect import bisect_left, bisect_right
from collections import defaultdict
from pathlib import Path
import re, sys, unittest, zipfile
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'src'))
from model import INPUT_ROWS

REF=re.compile(r"(?:(?:'([^']+)'|([A-Za-z_ÄÖÜäöüß][A-Za-z0-9_ÄÖÜäöüß]*))!)?(\$?[A-Z]{1,3}\$?\d+)(?::(\$?[A-Z]{1,3}\$?\d+))?")
STRINGS=re.compile(r'"(?:[^"]|"")*"')

class WorkbookTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.path=ROOT/'build/raw/Grauwasser_Bewertungstool.xlsx'
        cls.wb=load_workbook(cls.path)
    def test_no_external_files_or_macros(self):
        self.assertEqual(len(self.wb._external_links),0)
        with zipfile.ZipFile(self.path) as z:
            self.assertFalse(any('vbaProject' in n or n.startswith('xl/externalLinks/') for n in z.namelist()))
    def test_workflow_elements(self):
        self.assertEqual(len(self.wb.sheetnames),16)
        self.assertEqual(self.wb.active.title,'Dashboard')
        self.assertEqual(sum(len(s._charts) for s in self.wb),7)
        self.assertEqual(self.wb.calculation.calcMode,'auto')
        self.assertTrue(self.wb['Nutzertest'].tables)
    def test_inputs_are_editable_and_formula_cells_protected(self):
        for key,row in INPUT_ROWS.items(): self.assertFalse(self.wb['Eingaben'].cell(row,2).protection.locked,key)
        for ws in self.wb:
            self.assertTrue(ws.protection.sheet)
            for row in ws:
                for c in row:
                    if c.data_type=='f': self.assertTrue(c.protection.locked,(ws.title,c.coordinate))
    def test_all_formulas_have_valid_sheet_references(self):
        count=0
        for ws in self.wb:
            for row in ws:
                for c in row:
                    if c.data_type!='f': continue
                    count+=1; self.assertNotIn('#REF!',c.value)
                    self.assertNotIn('[',c.value,'Keine externen Workbook-Bezüge')
                    formula=STRINGS.sub('""',c.value)
                    for m in REF.finditer(formula):
                        name=m.group(1) or m.group(2)
                        if name: self.assertIn(name,self.wb.sheetnames,(ws.title,c.coordinate,name))
        self.assertGreater(count,90000)
    def test_dependency_graph_is_acyclic(self):
        nodes=[]; ids={}; columns=defaultdict(lambda:defaultdict(list))
        for ws in self.wb:
            for row in ws:
                for c in row:
                    if c.data_type=='f':
                        i=len(nodes); nodes.append((ws.title,c.coordinate,c.value)); ids[(ws.title,c.coordinate)]=i
                        columns[ws.title][c.column].append((c.row,i))
        for cs in columns.values():
            for values in cs.values(): values.sort()
        graph=[]
        for sheet,coord,formula in nodes:
            deps=set()
            for m in REF.finditer(STRINGS.sub('""',formula)):
                target=m.group(1) or m.group(2) or sheet
                a=m.group(3).replace('$',''); b=(m.group(4) or m.group(3)).replace('$','')
                if a==b:
                    if (target,a) in ids: deps.add(ids[(target,a)])
                else:
                    minc,minr,maxc,maxr=range_boundaries(a+':'+b)
                    for column in range(minc,maxc+1):
                        rows=columns[target].get(column,[])
                        lo=bisect_left(rows,(minr,-1)); hi=bisect_right(rows,(maxr,len(nodes)))
                        deps.update(i for _,i in rows[lo:hi])
            graph.append(deps)
        # Iterative DFS: eine ganze Jahres-Speicherreihe hat > 1000 Kanten Tiefe.
        states=bytearray(len(graph))
        for start in range(len(graph)):
            if states[start]: continue
            stack=[(start,iter(graph[start]))]; states[start]=1
            while stack:
                node,children=stack[-1]
                child=next(children,None)
                if child is None:
                    states[node]=2; stack.pop()
                elif states[child]==1:
                    self.fail('Zirkelbezug: '+str(nodes[node][:2])+' -> '+str(nodes[child][:2]))
                elif states[child]==0:
                    states[child]=1; stack.append((child,iter(graph[child])))
        self.assertEqual(sum(x==2 for x in states),len(nodes))

if __name__=='__main__': unittest.main(verbosity=2)
