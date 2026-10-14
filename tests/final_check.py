"""Letzte Nachprüfung: Diagramme bei anderem Horizont, ungültige Studienantworten."""
import json, subprocess, sys
from pathlib import Path
from tempfile import TemporaryDirectory
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'src')); sys.path.insert(0,str(ROOT/'tests'))
from model import Project,default_calendar
from engine_check import apply_project,check_file,compare

def main():
    results=[]
    with TemporaryDirectory(prefix='grauwasser-finalcheck-') as directory:
        directory=Path(directory); inputs=directory/'inputs'; outputs=directory/'outputs'; inputs.mkdir(); outputs.mkdir()
        for horizon in (5,35):
            p=Project(); p.horizon=horizon
            if horizon==5: p.replacement_year=p.replacement_cost=0
            w=load_workbook(ROOT/'build/raw/Grauwasser_Bewertungstool.xlsx'); apply_project(w,p,default_calendar(p.year))
            if horizon==5:
                s=w['Nutzertest']
                for row,uid,rating,t,errors in [(6,'P001',1,10,0),(7,'P002',3,20,2),(8,'P003','3','abc',1.5)]:
                    for c,v in [(1,uid),(6,rating),(4,t),(5,errors)]: s.cell(row,c,v)
                s['F9']=4 # Ohne ID: muss als nicht zugeordnete Zeile gemeldet werden.
            w.save(inputs/(str(horizon)+'.xlsx'))
        cmd=['soffice','-env:UserInstallation='+(directory/'profile').as_uri(),'--headless','--convert-to','xlsx','--outdir',str(outputs)]+[str(inputs/(str(n)+'.xlsx')) for n in (5,35)]
        proc=subprocess.run(cmd,capture_output=True,text=True,timeout=240); print(proc.stdout)
        if proc.returncode: raise RuntimeError(proc.stderr)
        for horizon in (5,35):
            p=Project(); p.horizon=horizon
            if horizon==5: p.replacement_year=p.replacement_cost=0
            path=outputs/(str(horizon)+'.xlsx')
            record=check_file(path,p,default_calendar(p.year),baseline=True)
            w=load_workbook(path,data_only=True)
            xs=[w['Dashboard'].cell(r,16).value for r in range(55,106)]
            compare(min(xs),0,'Diagrammstart'); compare(max(xs),horizon,'Diagrammhorizont')
            compare(w['Dashboard']['R105'].value,w['Berechnung']['B20'].value,'Endpunkt entspricht NPV')
            if horizon==5:
                out=w['Testauswertung']
                for c,v in [('B6',2),('C6',2),('K6',1),('B17',2),('B18',15),('B19',1),('B20',1),('B21',2),('B22',2)]: compare(out[c].value,v,'Auswertung '+c)
            formulas=load_workbook(path)
            for ch in formulas['Dashboard']._charts[1:]:
                assert ch.ser and ch.ser[0].xVal and ch.ser[0].yVal,'Diagrammreihe verloren'
                cache=ch.ser[0].xVal.numRef.numCache
                assert cache is not None and cache.ptCount==51,'Diagrammcache fehlt'
                compare(max(float(pt.v) for pt in cache.pt),horizon,'Gerenderte X-Achse')
            results.append({'horizon':horizon,'status':'passed',**record}); print('PASS finalcheck',horizon)
    (ROOT/'Validierung_Nachpruefung.json').write_text(json.dumps(results,ensure_ascii=False,indent=2)+'\n')

if __name__=='__main__': main()
