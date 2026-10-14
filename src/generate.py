"""Erzeugt eine eigenständige, makrofreie XLSX mit Excel-Formeln (OOXML Englisch)."""
from pathlib import Path
import argparse
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.chart import BarChart, LineChart, ScatterChart, Series, Reference
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter as col
from model import Project, INPUT_ROWS, default_calendar

NAVY='142D4E'; TEAL='007F86'; PALE='E8F4F5'; WHITE='FFFFFF'; GRAY='F2F5F8'; INK='22364A'
GREEN='D7EEDD'; YELLOW='FFF0BD'; RED='FAD9D8'; INPUT='E0F2F7'
EUR='#,##0.00 "€"'; VOL='#,##0.00 "m³"'; PCT='0.0%'; NUM='#,##0.00'; INT='0'
NAMES=['Dashboard','Eingaben','Nutzungsprofile','Kalender','Schwellenwerte','Wasserbilanz','Zahlungen','Berechnung','Szenarien','Variantenbilanz','VariantenCF','KPI_Katalog','Prüfungen','Nutzertest','Testauswertung','Anleitung']

def quoted(name): return "'"+name+"'"
def inp(key): return f"'Eingaben'!$B${INPUT_ROWS[key]}"
def calc(row): return f"'Berechnung'!$B${row}"
def setval(ws,addr,value,fmt=None,edit=False):
    c=ws[addr]; c.value=value
    c.font=Font(name='Aptos',size=11,color=TEAL if edit else INK)
    c.alignment=Alignment(vertical='center',wrap_text=True)
    if fmt: c.number_format=fmt
    if edit:
        c.fill=PatternFill('solid',fgColor=INPUT); c.protection=Protection(locked=False)
        c.border=Border(bottom=Side(style='thin',color=TEAL))
    return c

def dv(ws,cell,lo=0,hi=1,integer=False,choices=None):
    if choices:
        v=DataValidation(type='list',formula1='"'+','.join(choices)+'"',allow_blank=False)
    else:
        v=DataValidation(type='whole' if integer else 'decimal',operator='between',formula1=str(lo),formula2=str(hi),allow_blank=False)
    v.errorTitle='Eingabe prüfen'; v.error='Bitte einen gültigen Wert innerhalb des angegebenen Bereichs eingeben.'
    v.promptTitle='Veränderbare Eingabe'; v.prompt='Blaue Felder sind Eingaben. Einheiten und Hinweise beachten.'
    v.showErrorMessage=True; v.errorStyle='stop'; v.showInputMessage=True
    ws.add_data_validation(v); v.add(ws[cell])

def header(ws,title,subtitle,end=10):
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=end)
    c=ws.cell(1,1,title); c.fill=PatternFill('solid',fgColor=NAVY); c.font=Font(name='Aptos Display',size=20,bold=True,color=WHITE); c.alignment=Alignment(vertical='center')
    ws.row_dimensions[1].height=38
    ws.merge_cells(start_row=3,start_column=1,end_row=3,end_column=end)
    setval(ws,'A3',subtitle); ws.row_dimensions[3].height=34
    for i,(label,sheet) in enumerate([('Dashboard','Dashboard'),('Eingaben','Eingaben'),('Profile','Nutzungsprofile'),('Kalender','Kalender'),('Methodik','Anleitung')],1):
        c=ws.cell(2,i,label); c.hyperlink='#'+quoted(sheet)+'!A1'; c.font=Font(name='Aptos',size=10,bold=True,color=TEAL)
    ws.sheet_view.showGridLines=False; ws.freeze_panes='B5'
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='landscape'; ws.page_setup.paperSize=ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=1
    ws.print_options.horizontalCentered=True
    ws.oddFooter.center.text='Grauwasserbewertung | wissenschaftlicher Prototyp | &P / &N'
    ws.oddFooter.center.size=8
    ws.sheet_properties.tabColor=TEAL
    ws.page_margins.left=.25; ws.page_margins.right=.25
    for j in range(1,end+1): ws.column_dimensions[col(j)].width=16

def band(ws,row,text,end=4):
    ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=end)
    c=ws.cell(row,1,text); c.fill=PatternFill('solid',fgColor=TEAL); c.font=Font(name='Aptos',size=11,bold=True,color=WHITE)
    ws.row_dimensions[row].height=24

def table(ws,row,headers):
    for j,h in enumerate(headers,1):
        c=ws.cell(row,j,h); c.fill=PatternFill('solid',fgColor=NAVY); c.font=Font(name='Aptos',size=10,bold=True,color=WHITE); c.alignment=Alignment(wrap_text=True,vertical='center')
    ws.row_dimensions[row].height=32

def traffic(ws,region):
    for label,color in [('Grün',GREEN),('Gelb',YELLOW),('Rot',RED),('Nicht bewertbar',GRAY)]:
        ws.conditional_formatting.add(region,CellIsRule(operator='equal',formula=['"'+label+'"'],fill=PatternFill('solid',fgColor=color)))

def chart(ws,kind,title,data,cats,anchor,ytitle='',size=(15,8)):
    ch=BarChart() if kind=='bar' else ScatterChart() if kind=='scatter' else LineChart()
    ch.title=title; ch.style=13; ch.y_axis.title=ytitle; ch.x_axis.title=''
    if kind=='scatter':
        ch.scatterStyle='lineMarker'; ch.ser.append(Series(data,cats,title_from_data=True)); ch.x_axis.scaling.min=0
    else:
        ch.add_data(data,titles_from_data=True); ch.set_categories(cats)
    if kind=='scatter': ch.x_axis.title='Jahr'
    ch.width,ch.height=size
    for i,s in enumerate(ch.ser):
        color=[TEAL,NAVY,'7A9BB4','EAB05A'][i%4]
        s.graphicalProperties.solidFill=color; s.graphicalProperties.line.solidFill=color
        if kind!='bar': s.smooth=False
        if kind=='scatter':
            s.marker.symbol='circle'; s.marker.size=3
            s.marker.graphicalProperties.solidFill=color; s.marker.graphicalProperties.line.solidFill=color
    ws.add_chart(ch,anchor)

def make_inputs(wb,p):
    ws=wb['Eingaben']; header(ws,'Projekt & Anlage','Blaue Felder bearbeiten · RECHENBEISPIEL, nicht vermessene Projektwerte',4)
    ws.column_dimensions['A'].width=39; ws.column_dimensions['B'].width=48; ws.column_dimensions['C'].width=12; ws.column_dimensions['D'].width=83
    info={
    'name':('Projektname','–','Bezeichnung für Dashboard und Dokumentation.'),
    'building':('Gebäudetyp','–','Nutzergruppen separat in Nutzungsprofile erfassen.'),
    'location':('Standort','–','Gebührenbescheid und tatsächlichen Ferienkalender prüfen.'),
    'plant':('Geplante Anlage','–','Kapazität, Energie und Verluste mit Hersteller-/Messdaten belegen.'),
    'users':('Nutzer im Referenzprojekt','Personen','Informative Bezugsgröße; Profile dürfen Teil-/Fremdnutzer abbilden.'),
    'year':('Bilanzjahr','Jahr','Bei Jahreswechsel Kalenderflaggen kontrollieren; keine automatische Ferienerkennung.'),
    'horizon':('Betrachtungszeitraum','Jahre','1–50; Kalenderjahr wird für alle Betriebsjahre wiederholt.'),
    'other_water':('Sonstiger Trinkwasserbedarf','m³/a','Baseline außerhalb erfasster Betriebswassernutzungen; keine Doppelzählung.'),
    'yield_rate':('Aufbereitungsausbeute','%','Aufbereitet / verarbeitet. Kein garantierter Herstellerwert.'),
    'availability':('Betriebsverfügbarkeit','%','Pauschaler Tagesfaktor auf min(Rohwasser, Kapazität); kein Ausfallkalender.'),
    'capacity':('Rohwasser-Aufbereitungskapazität','m³/d','Maximale verarbeitbare Rohwassermenge je Tag vor Verfügbarkeitsfaktor.'),
    'storage':('Nutzbares Speichervolumen','m³','Modell: frischer Tagesrest für genau den Folgetag; alte Reste werden verworfen.'),
    'overlap':('Direkt nutzbarer Tagesanteil','%','Vereinfachte zeitliche Übereinstimmung; mit Stundenprofilen validieren.'),
    'storage_loss':('Verlust bei Tagesübertrag','%','Volumetrischer Verlust, keine hygienische Inaktivierungsrate.'),
    'hygiene':('Hygienische Nachweise','Status','Ungeklärt begrenzt Gesamtergebnis auf Gelb; Nicht erfüllt ergibt Rot.'),
    'technical':('Technische Nachweise','Status','Trennung, Nachspeisung, Betrieb und Auslegung projektbezogen prüfen.'),
    'legal':('Rechtliche Zulässigkeit','Status','Behördliche / Betreiberanforderungen separat nachweisen.'),
    'evidence':('Nachweise / Aktenverweise','–','Status ist Selbstauskunft; dieses Tool prüft keine Nachweisdokumente.'),
    'planning':('Planungskosten','€','Einheitlicher Preisstand und Brutto-/Nettobezug für alle Kosten.'),
    'equipment':('Anlagenkosten','€','Beschaffung; keine Montage doppelt erfassen.'),
    'installation':('Montagekosten','€','Einbau inklusive erforderlicher Leitungen erfassen.'),
    'extra_investment':('Sonstige Investitionskosten','€','Zusatzleistungen ohne Doppelzählung.'),
    'grant':('Fördermittel bei Investition','€','Jahr 0; höchstens Gesamtinvestition. In Varianten begrenzt auf Bruttoinvestition.'),
    'maintenance':('Wartungskosten erstes Jahr','€/a','Wiederkehrende Kosten; Preisentwicklung separat.'),
    'energy_intensity':('Spezifische Aufbereitungsenergie','kWh/m³','Bezug: verarbeitetes Rohwasser, auch bei später ungenutztem Wasser.'),
    'energy_price':('Strompreis erstes Jahr','€/kWh','Projektbezogener Tarif.'),
    'other_cost':('Sonstige Betriebskosten erstes Jahr','€/a','Labor, Verbrauchsmaterial, Betreiberaufwand etc.'),
    'water_price':('Vermiedener Trinkwasserpreis','€/m³','Nur mengenabhängig vermeidbarer Anteil; Grundgebühren ausschließen.'),
    'sewage_price':('Abwassergebühr','€/m³','Keine automatische Einsparung; Gebührenregelung projektspezifisch.'),
    'sewage_credit':('Anerkannter Abwasser-Einsparanteil','%','Standard 0 %. Positiver Wert erfordert dokumentierte Gebührenentlastung.'),
    'discount':('Kalkulationszins','%','Nominalzins zusammen mit nominalen Preissteigerungen; ≥ 0.'),
    'water_growth':('Trinkwasserpreissteigerung','%/a','Optionale nominale Szenarioannahme; Basisjahr t=1.'),
    'cost_growth':('Wartung / sonstige Preissteigerung','%/a','Auch für Ersatzinvestition; keine Inflationsprognose.'),
    'energy_growth':('Strompreissteigerung','%/a','Unabhängig veränderbar.'),
    'replacement_year':('Jahr der Ersatzinvestition','Jahr','0 nur ohne Ersatzkosten; sonst 1 bis Betrachtungszeitraum.'),
    'replacement_cost':('Ersatzinvestition im Preisstand Jahr 1','€','Ein Ersatzereignis; Kosten mit Betriebs-Kostensteigerung fortgeschrieben.'),
    'residual':('Restwert am Betrachtungsende','€','Optionaler nominaler Endwert; standardmäßig 0. Keine automatische Wertermittlung.'),
    'sewage_growth':('Abwassergebührensteigerung','%/a','Nur auf bestätigten mengenabhängigen Einsparanteil angewandt.'),
    'billing_evidence':('Beleg der Gebührenentlastung','–','Bei positiver Anerkennung voreingestellten Text durch konkreten Beleg ersetzen.'),
    'data_status':('Datenstatus','Status','Beispieldaten sind keine Entscheidungsgrundlage für das reale Projekt.')}
    for r,t in [(4,'01  Projekt'),(14,'02  Anlage, Speicher & Nachweise'),(26,'03  Wirtschaftlichkeit')]: band(ws,r,t)
    limits={'other_water':(0,1e7,False),'capacity':(0,1e5,False),'storage':(0,1e5,False),'users':(0,1000000,True),'year':(2020,2100,True),'horizon':(1,50,True),'replacement_year':(0,50,True)}
    for key,row in INPUT_ROWS.items():
        label,unit,note=info[key]; setval(ws,f'A{row}',label); setval(ws,f'B{row}',getattr(p,key),edit=True); setval(ws,f'C{row}',unit); setval(ws,f'D{row}',note)
        ws.row_dimensions[row].height=36 if isinstance(getattr(p,key),str) else 30
        if key in ('hygiene','technical','legal'): dv(ws,f'B{row}',choices=['Ungeklärt','Nachgewiesen','Nicht erfüllt'])
        elif key=='data_status': dv(ws,f'B{row}',choices=['Beispieldaten','Projektdaten'])
        elif not isinstance(getattr(p,key),str):
            if key in limits: lo,hi,integer=limits[key]
            elif key in ('yield_rate','availability','overlap','storage_loss','sewage_credit','discount'): lo,hi,integer=0,1,False
            elif key.endswith('growth'): lo,hi,integer=-.5,1,False
            else: lo,hi,integer=0,1e9,False
            dv(ws,f'B{row}',lo,hi,integer)
            ws[f'B{row}'].number_format=PCT if key.endswith('growth') or key in ('yield_rate','availability','overlap','storage_loss','sewage_credit','discount') else INT if integer else NUM
    ws.print_area='A1:D48'

def make_profiles(wb,p):
    ws=wb['Nutzungsprofile']; header(ws,'Nutzergruppen & Mengen','Tage sind Jahres-Nutzungstage je Profil; Verteilung auf aktive Kalendertage der Gruppe. Quellen fachlich separat freigeben.',10)
    headers=['Quelle / Nutzung','Kalendergruppe','Nutzer','Nutzungen je Tag','Liter je Nutzung','Tage pro Jahr','Jahresmenge m³','Aktive Kalendertage','Prüfung','Hinweise / Quelle']
    for start,profiles in [(6,p.sources),(17,p.demands)]:
        table(ws,start-1,headers)
        for r,item in enumerate(profiles,start):
            for j,v in enumerate([item.name,item.group,item.users,item.frequency,item.liters,item.days],1):
                setval(ws,f'{col(j)}{r}',v,INT if j in (3,6) else NUM if j in (4,5) else None,edit=True)
            dv(ws,f'B{r}',choices=['Schulbetrieb','Halle','Individuell']); dv(ws,f'C{r}',0,1e6); dv(ws,f'D{r}',0,100); dv(ws,f'E{r}',0,1000); dv(ws,f'F{r}',0,366,True)
            setval(ws,f'G{r}',f'=IFERROR(C{r}*D{r}*E{r}*F{r}/1000,0)',NUM)
            setval(ws,f'H{r}',f'=IF(B{r}="Schulbetrieb",SUM(Kalender!$C$5:$C$370),IF(B{r}="Halle",SUM(Kalender!$D$5:$D$370),IF(B{r}="Individuell",SUM(Kalender!$E$5:$E$370),0)))',INT)
            checks=[f'LEN(TRIM(A{r}))>0',f'OR(B{r}="Schulbetrieb",B{r}="Halle",B{r}="Individuell")',f'COUNT(C{r}:F{r})=4',f'C{r}>=0',f'C{r}<=1000000',f'D{r}>=0',f'D{r}<=100',f'E{r}>=0',f'E{r}<=1000',f'F{r}>=0',f'F{r}<=366',f'MOD(F{r},1)=0',f'OR(G{r}=0,F{r}<=H{r})']
            setval(ws,f'I{r}','=IFERROR(IF(AND('+','.join(checks)+'),"OK","FEHLER"),"FEHLER")')
            setval(ws,f'J{r}','Fiktiver Rechenansatz; Messung / Planungsnachweis ergänzen.',edit=True)
            ws.row_dimensions[r].height=38
    for row,text in [(13,'Grauwasserquellen: Küchenabwasser und belastete Quellen nicht ungeprüft als zulässig ansetzen.'),(24,'Tagesverteilung = Jahresmenge / aktive Gruppentage. Weniger Profiltage reduzieren die Tagesintensität gleichmäßig.'),(25,'Schule und Halle dürfen verschiedene Nutzerzahlen / Tage haben. Anzahl im Projektkopf ist nur informativ.'),(26,'Nutzungstage größer als aktive Gruppentage bei positivem Volumen sind ungültig. Kalender entsprechend ergänzen.')]:
        ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=10); setval(ws,f'A{row}',text); ws.row_dimensions[row].height=30
    for c,w in {'A':34,'B':20,'C':12,'D':16,'E':16,'F':14,'G':18,'H':19,'I':13,'J':50}.items(): ws.column_dimensions[c].width=w
    ws.auto_filter.ref='A5:J11'; ws.print_area='A1:J26'

def make_calendar(wb,p):
    ws=wb['Kalender']; header(ws,'Tageskalender','0 = inaktiv · 1 = aktiv. Fiktiv: 190 Schul-, 280 Hallentage; Ferien/Feiertage selbst abgleichen. Hallentage enthalten Wochenenden.',7)
    table(ws,4,['Datum','Wochentag','Schulbetrieb 0/1','Halle 0/1','Individuell 0/1','Notiz / Datengrundlage','Prüfung'])
    schedule=default_calendar(p.year)
    for r in range(5,371):
        i=r-5; flags=schedule[i] if i<len(schedule) else (0,0,0)
        setval(ws,f'A{r}',f'=IFERROR(IF(YEAR(DATE({inp("year")},1,1)+{i})={inp("year")},DATE({inp("year")},1,1)+{i},""),"")','dd.mm.yyyy')
        setval(ws,f'B{r}',f'=IF(A{r}="","",CHOOSE(WEEKDAY(A{r},2),"Mo","Di","Mi","Do","Fr","Sa","So"))')
        for j,v in enumerate(flags,3): setval(ws,f'{col(j)}{r}',v,INT,True); dv(ws,f'{col(j)}{r}',0,1,True)
        setval(ws,f'F{r}','Fiktives Kalenderprofil – zu ersetzen' if i==0 else '',edit=True)
        setval(ws,f'G{r}',f'=IFERROR(IF(AND(COUNT(C{r}:E{r})=3,OR(C{r}=0,C{r}=1),OR(D{r}=0,D{r}=1),OR(E{r}=0,E{r}=1),OR(A{r}<>"",SUM(C{r}:E{r})=0)),"OK","FEHLER"),"FEHLER")')
    ws.column_dimensions['F'].width=55; ws.auto_filter.ref='A4:G370'; ws.freeze_panes='C5'; ws.print_area='A1:G40'

def make_thresholds(wb,p):
    ws=wb['Schwellenwerte']; header(ws,'Ampelmodell & Schwellen','Alle Grenzwerte sind vorläufige Modellannahmen, keine normativen oder empirisch validierten Eignungsgrenzen.',5)
    table(ws,4,['Kriterium','Schwelle','Einheit','Status / Interpretation','Begründung'])
    rows=[(5,'Ökonomie Grün ab',p.economic_green,'NPV / Nettoinvestition','Grün ab Kapitalwert ≥ 0 bei Basisschwelle 0 %.'),
          (6,'Ökonomie Gelb ab',p.economic_yellow,'NPV / Nettoinvestition','Verlust bis 20 % des Kapitaleinsatzes als vorläufiger Prüfbereich.'),
          (9,'Ressourcen Grün ab',p.resource_green,'Anteil Baseline-Trinkwasser','Vorläufig: 15 % der Baseline aus Betriebsbedarf + sonstigem Bedarf.'),
          (10,'Ressourcen Gelb ab',p.resource_yellow,'Anteil Baseline-Trinkwasser','Vorläufig: 5 %; keine Literaturgrenze.')]
    for r,label,value,unit,note in rows:
        setval(ws,f'A{r}',label); setval(ws,f'B{r}',value,PCT,True); dv(ws,f'B{r}',-1 if r<8 else 0,1)
        setval(ws,f'C{r}',unit); setval(ws,f'D{r}','Vorläufig · zu validieren'); setval(ws,f'E{r}',note); ws.row_dimensions[r].height=45
    notes=[
    'Aggregation: schlechteste Kategorie entscheidet (konservatives Maximum). Keine kompensatorische gewichtete Summe.',
    'Grün / Gelb / Rot entsprechen günstiger / bedingter / ungünstiger Eignung innerhalb dieses Modells.',
    'Nachweise: Nicht erfüllt → Rot; Ungeklärt → mindestens Gelb + Prüfbedarf; alle Nachgewiesen → keine Begrenzung.',
    'Fehlende oder ungültige Pflichtwerte → Nicht bewertbar; kein Ersatz durch Nullwerte im Dashboard.',
    'Keine Wassereinsparung → Ressourcen Rot unabhängig von Schwellen. Kein Bedarf ist kein positiver Nutzen.',
    'Bei Nettoinvestition = 0: NPV ≥ 0 → ökonomisch Grün, NPV < 0 → Rot; Quotient ist nicht definiert.',
    'Technisch/hygienisch/rechtlich sind Nachweisstatus, keine vom Tool erteilte Freigabe. Selbstauskunft dokumentieren.',
    'Gewichtete Bewertung wurde verworfen: Kompensation verdeckt Konflikte, korrelierte KPIs würden doppelt zählen.',
    'Konservative Aggregation macht Zielkonflikte sichtbar; alternative Gewichte erst mit Beteiligten validieren.',
    'Wissenschaftliche Validierung: Schwellen begründen, Expertenreview und Sensitivität durchführen; Nutzertest separat.']
    for r,text in enumerate(notes,13): ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=5); setval(ws,f'A{r}',text); ws.row_dimensions[r].height=34
    ws.column_dimensions['A'].width=33; ws.column_dimensions['C'].width=29; ws.column_dimensions['D'].width=30; ws.column_dimensions['E'].width=67; ws.print_area='A1:E22'

def make_checks(wb,p):
    ws=wb['Prüfungen']; header(ws,'Eingabe- & Plausibilitätsprüfung','FEHLER sperrt Projektbewertung; Hinweise stehen zusätzlich im Dashboard.',5)
    setval(ws,'A4','Anzahl Pflichtfehler'); table(ws,6,['Parameter / Zusammenhang','Bedingung','Fundstelle','Ergebnis','Aktion bei Fehler'])
    r=7
    def add(label,condition,location,note):
        nonlocal r
        setval(ws,f'A{r}',label); setval(ws,f'B{r}',condition); setval(ws,f'C{r}',location)
        setval(ws,f'D{r}',f'=IFERROR(IF({condition},"OK","FEHLER"),"FEHLER")'); setval(ws,f'E{r}',note); r+=1
    integers={'users','year','horizon','replacement_year'}
    bounds={'users':(0,1e6),'year':(2020,2100),'horizon':(1,50),'replacement_year':(0,50),'other_water':(0,1e7),'capacity':(0,1e5),'storage':(0,1e5)}
    for key,row in INPUT_ROWS.items():
        ref=inp(key); value=getattr(p,key)
        if key in ('hygiene','technical','legal'): cond=f'OR({ref}="Ungeklärt",{ref}="Nachgewiesen",{ref}="Nicht erfüllt")'
        elif key=='data_status': cond=f'OR({ref}="Beispieldaten",{ref}="Projektdaten")'
        elif isinstance(value,str): cond=f'AND(ISTEXT({ref}),LEN(TRIM({ref}))>0)'
        else:
            if key in bounds: lo,hi=bounds[key]
            elif key in ('yield_rate','availability','overlap','storage_loss','sewage_credit','discount'): lo,hi=0,1
            elif key.endswith('growth'): lo,hi=-.5,1
            else: lo,hi=0,1e9
            cond=f'AND(ISNUMBER({ref}),{ref}>={lo},{ref}<={hi}'+(f',MOD({ref},1)=0' if key in integers else '')+')'
        add(key,cond,f'Eingaben!B{row}','Pflichtfeld, Datentyp und Wertebereich prüfen.')
    add('Förderung ≤ Bruttoinvestition',f'{inp("grant")}<=SUM(Eingaben!B27:B30)','Eingaben!B31','Förderung darf Kosten nicht übersteigen.')
    add('Ersatzjahr innerhalb Horizont',f'{inp("replacement_year")}<={inp("horizon")}','Eingaben!B43','Jahr 0 nur ohne Ersatzkosten.')
    add('Ersatzkosten benötigen Jahr',f'OR({inp("replacement_cost")}=0,{inp("replacement_year")}>0)','Eingaben!B43:B44','Jahr 1 bis Betrachtungsende wählen.')
    add('Abwasserentlastung belegt',f'OR({inp("sewage_credit")}=0,AND(LEN(TRIM({inp("billing_evidence")}))>0,LEFT({inp("billing_evidence")},LEN("Keine bestätigte"))<>"Keine bestätigte"))','Eingaben!B38/B47','Dokumentierten Gebührenbescheid ergänzen; Standardtext ersetzen.')
    add('Ressourcenschwellen geordnet','AND(ISNUMBER(Schwellenwerte!B9),ISNUMBER(Schwellenwerte!B10),Schwellenwerte!B10>=0,Schwellenwerte!B9<=1,Schwellenwerte!B10<=Schwellenwerte!B9)','Schwellenwerte!B9:B10','0 ≤ Gelb ≤ Grün ≤ 100 %.')
    add('Ökonomieschwellen geordnet','AND(ISNUMBER(Schwellenwerte!B5),ISNUMBER(Schwellenwerte!B6),Schwellenwerte!B6>=-1,Schwellenwerte!B5<=1,Schwellenwerte!B6<=Schwellenwerte!B5)','Schwellenwerte!B5:B6','−100 % ≤ Gelb ≤ Grün ≤ 100 %.')
    add('Profile gültig','AND(COUNTIF(Nutzungsprofile!I6:I11,"FEHLER")=0,COUNTIF(Nutzungsprofile!I17:I22,"FEHLER")=0)','Nutzungsprofile!I','Numerische Felder, Gruppe und aktive Tage prüfen.')
    add('Kalender gültig','COUNTIF(Kalender!G5:G370,"FEHLER")=0','Kalender!G','Nur 0/1; in nicht existierendem 366. Tag alle Flaggen 0.')
    setval(ws,'B4',f'=COUNTIF(D7:D{r-1},"FEHLER")',INT)
    for c,w in {'A':39,'B':75,'C':28,'D':15,'E':60}.items(): ws.column_dimensions[c].width=w
    ws.conditional_formatting.add(f'D7:D{r-1}',CellIsRule(operator='equal',formula=['"FEHLER"'],fill=PatternFill('solid',fgColor=RED)))
    ws.auto_filter.ref=f'A6:E{r-1}'; ws.print_area=f'A1:E{r-1}'

def daily_profile(start,end,calrow):
    terms=[]
    for group,column in [('Schulbetrieb','C'),('Halle','D'),('Individuell','E')]:
        terms.append(f'IFERROR(SUMIF(Nutzungsprofile!$B${start}:$B${end},"{group}",Nutzungsprofile!$G${start}:$G${end})*Kalender!{column}{calrow}/SUM(Kalender!${column}$5:${column}$370),0)')
    return '('+'+'.join(terms)+')'

def make_water(wb):
    ws=wb['Wasserbilanz']; header(ws,'Wasserbilanz · 366 Tageszeilen','m³ je Tag · Speicher aus Vortag zuerst nutzen, dann direkte Deckung; alte Reste spätestens am Tagesende verwerfen.',17)
    table(ws,4,['Datum','Schule','Halle','Rohgrauwasser','Betriebsbedarf','Verarbeitet','Aufbereitet','Altbestand nach Verlust','Verbrauch Alt','Direktverbrauch','Endbestand frisch','Wiederverwendung','Nachspeisung','Überschuss frisch','Speicherverlust / Alt','Bilanzrest','Tank Vortag'])
    for r in range(5,371):
        formulas={
        'A':f'Kalender!A{r}','B':f'Kalender!C{r}','C':f'Kalender!D{r}',
        'D':f'IF(Prüfungen!$B$4=0,{daily_profile(6,11,r)},0)',
        'E':f'IF(Prüfungen!$B$4=0,{daily_profile(17,22,r)},0)',
        'F':f'IFERROR(MIN(D{r},{inp("capacity")})*{inp("availability")},0)',
        'G':f'IFERROR(F{r}*{inp("yield_rate")},0)',
        'Q':'0' if r==5 else f'IF(A{r}="",0,K{r-1})',
        'H':f'IFERROR(Q{r}*(1-{inp("storage_loss")}),0)',
        'I':f'MIN(H{r},E{r})',
        'J':f'IFERROR(MIN(G{r}*{inp("overlap")},MAX(0,E{r}-I{r})),0)',
        'K':f'IFERROR(MIN({inp("storage")},MAX(0,G{r}-J{r})),0)',
        'L':f'I{r}+J{r}','M':f'MAX(0,E{r}-L{r})','N':f'MAX(0,G{r}-J{r}-K{r})',
        'O':f'Q{r}-I{r}',
        'P':f'G{r}+Q{r}-L{r}-N{r}-O{r}-K{r}'}
        for c,f in formulas.items(): setval(ws,f'{c}{r}','='+f,'dd.mm.yyyy' if c=='A' else NUM)
    # Tag 366 wird bei Nicht-Schaltjahren nicht gerechnet. Restbestand wird über den letzten echten Tag bilanziert.
    ws.auto_filter.ref='A4:Q370'; ws.print_area='A1:Q25'

def make_payments(wb):
    ws=wb['Zahlungen']; header(ws,'Zahlungsreihe & Kapitalwert','Nominale Zahlungen am Jahresende; Jahr 0 = Nettoinvestition. Jahreswasserprofil wiederholt sich ohne Wachstum der Nutzerzahl.',13)
    table(ws,6,['Jahr','Einsparmenge m³','Vermiedene Kosten €','Betriebskosten €','Ersatzinvestition €','Restwert €','Cashflow €','Diskontfaktor','Barwert CF €','Kumuliert diskontiert €','Barwert Kosten €','Barwert Einsparmenge m³','Erste Rückgewinnung Jahr'])
    for r in range(7,58):
        t=r-7; setval(ws,f'A{r}',t,INT)
        if t==0: formulas={'B':'0','C':'0','D':'0','E':calc(14),'F':'0','G':f'-E{r}','H':'1','I':f'G{r}','J':f'I{r}','K':f'E{r}','L':'0','M':'""'}
        else:
            active=f'AND(A{r}<={inp("horizon")},Prüfungen!$B$4=0)'
            formulas={'B':f'IF({active},{calc(10)},0)',
              'C':f'IFERROR(IF({active},B{r}*({inp("water_price")}*(1+{inp("water_growth")})^(A{r}-1)+{inp("sewage_price")}*{inp("sewage_credit")}*(1+{inp("sewage_growth")})^(A{r}-1)),0),0)',
              'D':f'IFERROR(IF({active},({inp("maintenance")}+{inp("other_cost")})*(1+{inp("cost_growth")})^(A{r}-1)+{calc(15)}*{inp("energy_price")}*(1+{inp("energy_growth")})^(A{r}-1),0),0)',
              'E':f'IFERROR(IF(AND({active},A{r}={inp("replacement_year")}),{inp("replacement_cost")}*(1+{inp("cost_growth")})^(A{r}-1),0),0)',
              'F':f'IFERROR(IF(AND({active},A{r}={inp("horizon")}),{inp("residual")},0),0)',
              'G':f'C{r}-D{r}-E{r}+F{r}',
              'H':f'IFERROR(IF({active},1/(1+{inp("discount")})^A{r},0),0)',
              'I':f'G{r}*H{r}','J':f'J{r-1}+I{r}',
              'K':f'(D{r}+E{r}-F{r})*H{r}','L':f'B{r}*H{r}',
              'M':f'IF(AND({active},{calc(14)}>0,J{r-1}<0,J{r}>=0,I{r}>0),A{r}-1-J{r-1}/I{r},"")'}
        for c,f in formulas.items(): setval(ws,f'{c}{r}','='+f,NUM)
    ws.print_area='A1:M28'

def label(score): return f'IF({score}=1,"Grün",IF({score}=2,"Gelb","Rot"))'

def make_calc(wb):
    ws=wb['Berechnung']; header(ws,'Kennzahlen · nachvollziehbarer Rechenkern','Ungültige Eingaben werden im Dashboard gesperrt. Technische Hilfswerte sind keine Ergebnisse bei Pflichtfehlern.',4)
    table(ws,4,['Kennzahl','Wert','Einheit','Definition / Bezug'])
    vals={
    5:('Pflichtfehler','Prüfungen!B4','Anzahl','Muss 0 sein.'),
    6:('Jährlicher Grauwasseranfall','SUM(Wasserbilanz!D5:D370)','m³/a','Rohgrauwasser aller Quellen.'),
    7:('Verarbeitetes Rohgrauwasser','SUM(Wasserbilanz!F5:F370)','m³/a','Kapazität und Verfügbarkeit begrenzen Verarbeitung.'),
    8:('Aufbereitete Wassermenge','SUM(Wasserbilanz!G5:G370)','m³/a','Verarbeitet × Ausbeute.'),
    9:('Betriebswasserbedarf','SUM(Wasserbilanz!E5:E370)','m³/a','Bedarf aller zulässigen Anwendungen.'),
    10:('Nutzbare Menge / Trinkwassereinsparung','SUM(Wasserbilanz!L5:L370)','m³/a','Altverbrauch + Direktverbrauch; ohne Nachspeisung.'),
    11:('Erforderliche Trinkwassernachspeisung','SUM(Wasserbilanz!M5:M370)','m³/a','Bedarf − Wiederverwendung.'),
    12:('Nicht genutztes aufbereitetes Wasser','B8-B10','m³/a','Überschuss + Speicherverluste + Endbestand.'),
    13:('Gesamtinvestition','IFERROR(SUM(Eingaben!B27:B30),0)','€','Planung + Anlage + Montage + Sonstiges.'),
    14:('Nettoinvestition',f'IFERROR(MAX(0,B13-{inp("grant")}),0)','€','Gesamtinvestition − Förderung.'),
    15:('Jährlicher Strombedarf',f'IFERROR(B7*{inp("energy_intensity")},0)','kWh/a','Energie auf verarbeitete Rohmenge, kein LCA-Nachweis.'),
    16:('Betriebskosten Jahr 1',f'IFERROR({inp("maintenance")}+{inp("other_cost")}+B15*{inp("energy_price")},0)','€/a','Wartung, Sonstiges, Energie.'),
    17:('Vermiedene Kosten Jahr 1',f'IFERROR(B10*({inp("water_price")}+{inp("sewage_price")}*{inp("sewage_credit")}),0)','€/a','Nur tatsächlich vermiedene variable Gebühren.'),
    18:('Netto-Cashflow Jahr 1','B17-B16','€/a','Vor Ersatzinvestition.'),
    19:('Statische Amortisation','IF(B14=0,"Keine Nettoinvestition",IF(B18<=0,"Nicht erreichbar",B14/B18))','Jahre','Konstante Jahr-1-Nettoeinsparung; ohne Zins/Ersatz.'),
    20:('Kapitalwert','SUM(Zahlungen!I7:I57)','€','Summe aller diskontierten Cashflows incl. Jahr 0.'),
    21:('Erste dynamische Amortisation','IF(B14=0,"Keine Nettoinvestition",IF(COUNT(Zahlungen!M8:M57)=0,"Im Zeitraum nicht erreicht",MIN(Zahlungen!M8:M57)))','Jahre','Erste Rückgewinnung; interpoliert. Stabilität separat prüfen.'),
    22:('Trinkwassereinsparquote',f'IFERROR(B10/(B9+{inp("other_water")}),0)','%','Trinkwasser-Baseline = Betriebsbedarf + sonstiger Bedarf.'),
    23:('Deckungsgrad','IFERROR(B10/B9,0)','%','Wiederverwendung / Betriebsbedarf.'),
    24:('Grauwassernutzungsgrad','IFERROR(B10/B6,0)','%','Wiederverwendung / Rohgrauwasseranfall.'),
    25:('Spezifische Lebenszykluskosten','IF(SUM(Zahlungen!L7:L57)>0,SUM(Zahlungen!K7:K57)/SUM(Zahlungen!L7:L57),"Keine Einsparmenge")','€/m³','Barwert Kosten / Barwert eingespartes Volumen.'),
    26:('Frischer Überschuss','SUM(Wasserbilanz!N5:N370)','m³/a','Nicht speicherbarer frischer Tagesrest.'),
    27:('Nachweis-Sperre',f'IF(OR({inp("hygiene")}="Nicht erfüllt",{inp("technical")}="Nicht erfüllt",{inp("legal")}="Nicht erfüllt"),3,IF(OR({inp("hygiene")}="Ungeklärt",{inp("technical")}="Ungeklärt",{inp("legal")}="Ungeklärt"),2,1))','1/2/3','Maximum mit Kategorien, niemals kompensiert.'),
    28:('Ökonomischer Score','IF(B14=0,IF(B20>=0,1,3),IF(B20/B14>=Schwellenwerte!B5,1,IF(B20/B14>=Schwellenwerte!B6,2,3)))','1/2/3','Kapitalwert relativ zur Nettoinvestition.'),
    29:('Ressourcen-Score','IF(B10<=0,3,IF(B22>=Schwellenwerte!B9,1,IF(B22>=Schwellenwerte!B10,2,3)))','1/2/3','Nur Einsparquote wird bewertet.'),
    30:('Gesamt-Score','MAX(B27:B29)','1/2/3','Schlechteste Kategorie einschließlich Nachweisen.'),
    31:('Ökonomische Ampel',f'IF(B5>0,"Nicht bewertbar",{label("B28")})','Status','Günstig / bedingt / ungünstig.'),
    32:('Ressourcenampel',f'IF(B5>0,"Nicht bewertbar",{label("B29")})','Status','Keine Aussage über gesamtheitliche Umweltbilanz.'),
    33:('Gesamtbewertung',f'IF(B5>0,"Nicht bewertbar",{label("B30")})','Status','Keine behördliche oder normative Freigabe.'),
    34:('Nachweisstatus','IF(B27=3,"Mindestens ein Nachweis nicht erfüllt",IF(B27=2,"PRÜFBEDARF: Nachweise ungeklärt","Alle Nachweise als erbracht eingetragen"))','Status','Selbstauskunft, Belege separat prüfen.'),
    35:('Jahresbilanzrest','B8-B10-B26-B38-B39','m³/a','Aufbereitet = Nutzung + Überschuss + Verlust + Endbestand.'),
    36:('Rohwasser-Kontrollrest','SUM(Nutzungsprofile!G6:G11)-B6','m³/a','Profil-Jahresmenge minus Tagesintegration.'),
    37:('Dynamische Rückgewinnung stabil?',f'IF(ISNUMBER(B21),IF(SUMPRODUCT((Zahlungen!A8:A57>=ROUNDUP(B21,0))*(Zahlungen!A8:A57<={inp("horizon")} )*(Zahlungen!J8:J57<0))=0,"Stabil bis Horizont","Spätere Unterdeckung"),"Keine positive Rückgewinnung")','Status','Ersatzinvestition kann erste Amortisation rückgängig machen.'),
    38:('Speicherverluste / verworfene Altmenge','SUM(Wasserbilanz!O5:O370)','m³/a','Übertragsverluste + ungenutzte Vortragsmenge.'),
    39:('Endbestand Bilanzjahr',f'IFERROR(INDEX(Wasserbilanz!K5:K370,DATE({inp("year")}+1,1,1)-DATE({inp("year")},1,1)),0)','m³','Jahresende enthält keine im Jahr verbuchte Einsparung.'),
    40:('Kritische Hinweise',f'IF({inp("data_status")}="Beispieldaten","BEISPIELDATEN · ","")&IF(B5>0,"Eingaben korrigieren (Prüfungen). · ","")&IF(B27>1,B34&" · ","")&IF(B10=0,"Keine Wassereinsparung. · ","")&IF(B18<=0,"Keine positive Nettoeinsparung in Jahr 1. · ","")&IF(AND(ISNUMBER(B19),B19>{inp("horizon")}),"Statische Amortisation außerhalb Horizont. · ","")&"Schwellen vorläufig; Tagesprofil und Quellen projektbezogen validieren."','Text','Kein Freigabenachweis.'),
    41:('Bewertungsstatus',f'IF(B5>0,"Nicht bewertbar",IF({inp("data_status")}="Beispieldaten","Modellergebnis mit Beispieldaten",IF(B27>1,"Modellergebnis mit Prüfbedarf","Modellergebnis – Nachweise eingetragen")))','Status','Datenqualität und fachliche Freigabe sind getrennt.'),
    42:('Aufbereitungsverlust','B7-B8','m³/a','Verarbeitet minus aufbereitet.'),
    43:('Nicht verarbeitetes Rohwasser','B6-B7','m³/a','Kapazitäts-/Verfügbarkeitsbegrenzung.')}
    for r,(name,f,unit,note) in vals.items():
        setval(ws,f'A{r}',name); setval(ws,f'B{r}','=IFERROR('+f+',"Nicht bewertbar")' if r in (28,29,30) else '='+f,PCT if r in (22,23,24) else NUM); setval(ws,f'C{r}',unit); setval(ws,f'D{r}',note); ws.row_dimensions[r].height=34
    ws.column_dimensions['A'].width=48; ws.column_dimensions['B'].width=53; ws.column_dimensions['C'].width=14; ws.column_dimensions['D'].width=72
    traffic(ws,'B31:B33'); ws.print_area='A1:D43'

def make_scenarios(wb):
    ws=wb['Szenarien']; header(ws,'Szenarien & Sensitivität','Jede Zeile rechnet die gesamte Tagesbilanz erneut. Faktoren sind frei veränderbare Annahmen; Zinsänderung in Prozentpunkten.',16)
    table(ws,7,['Variante','Investitionsfaktor','Wasserpreisfaktor','Bedarfsfaktor','Grauwasserfaktor','Betriebskostenfaktor','Zins Δ','Prüfung','Nutzung m³/a','Einsparquote','CF Jahr 1 €','Kapitalwert €','Statisch Jahre','Ökonomie','Ressourcen','Gesamt'])
    cases=[(8,'Konservativ',(1.2,.9,.9,.8,1.2,.01)),(9,'Basis',(1,1,1,1,1,0)),(10,'Optimistisch',(.8,1.1,1.1,1.2,.8,-.01))]
    labels=[('Investition',0),('Trinkwasserpreis',1),('Betriebswasserbedarf',2),('Grauwasseranfall',3),('Betriebskosten',4),('Kalkulationszins',5)]
    for i,(name,idx) in enumerate(labels):
        for j,direction in enumerate((-1,1)):
            f=[1,1,1,1,1,0]; f[idx]=direction*.01 if idx==5 else 1+direction*.2
            cases.append((14+2*i+j,name+(' −1 %-Punkt' if idx==5 and direction<0 else ' +1 %-Punkt' if idx==5 else ' −20 %' if direction<0 else ' +20 %'),tuple(f)))
    vws=wb['Variantenbilanz']; header(vws,'Tagesbilanzen aller Varianten','Berechnung, keine Eingaben. Je Variante 366 Tageszeilen; gleiche Kalender- und Speicherregeln wie im Basismodell.',14)
    table(vws,4,['Variante','Tag','Rohmenge','Bedarf','Verarbeitet','Aufbereitet','Alt nach Verlust','Altverbrauch','Direktverbrauch','Endbestand','Nutzung','Nachspeisung','Überschuss','Speicherverluste'])
    cf=wb['VariantenCF']; header(cf,'Zahlungsreihen aller Varianten','Berechnung, keine Eingaben. Faktoren gelten auch für Ersatzinvestition und Restwert; Förderung bleibt absolut gedeckelt.',11)
    table(cf,4,['Variante','Jahr','Nutzen','Betriebskosten','Investition / Ersatz','Restwert','Cashflow','Diskontfaktor','Barwert CF','Barwert Kosten','Barwert Menge'])
    for i,(r,name,factors) in enumerate(cases):
        setval(ws,f'A{r}',name)
        for j,v in enumerate(factors,2): setval(ws,f'{col(j)}{r}',v,PCT if j==7 else '0.00"×"',True); dv(ws,f'{col(j)}{r}',-1 if j==7 else 0,1 if j==7 else 10)
        setval(ws,f'H{r}',f'=IFERROR(IF(AND(Prüfungen!$B$4=0,COUNT(B{r}:G{r})=6,MIN(B{r}:F{r})>=0,MAX(B{r}:F{r})<=10,{inp("discount")}+G{r}>=0,{inp("discount")}+G{r}<=1),"OK","FEHLER"),"FEHLER")')
        start=5+i*366; end=start+365
        for x in range(start,end+1):
            calrow=5+x-start
            sf=f"'Szenarien'!$E${r}"; df=f"'Szenarien'!$D${r}"; valid=f"'Szenarien'!$H${r}=\"OK\""
            forms={
             'A':f'"{name}"' if x==start else '""','B':str(x-start+1),
             'C':f'IF({valid},{daily_profile(6,11,calrow)}*{sf},0)',
             'D':f'IF({valid},{daily_profile(17,22,calrow)}*{df},0)',
             'E':f'IFERROR(MIN(C{x},{inp("capacity")})*{inp("availability")},0)',
             'F':f'IFERROR(E{x}*{inp("yield_rate")},0)',
             'G':'0' if x==start else f'IFERROR(IF(Kalender!A{calrow}="",0,J{x-1}*(1-{inp("storage_loss")})),0)',
             'H':f'MIN(G{x},D{x})','I':f'IFERROR(MIN(F{x}*{inp("overlap")},MAX(0,D{x}-H{x})),0)',
             'J':f'IFERROR(MIN({inp("storage")},MAX(0,F{x}-I{x})),0)',
             'K':f'H{x}+I{x}','L':f'MAX(0,D{x}-K{x})','M':f'MAX(0,F{x}-I{x}-J{x})',
             'N':'0' if x==start else f'IF(Kalender!A{calrow}="",0,J{x-1}-H{x})'}
            for c,f in forms.items(): setval(vws,f'{c}{x}','='+f,NUM)
        used=f'SUM(Variantenbilanz!K{start}:K{end})'; demand=f'SUM(Variantenbilanz!D{start}:D{end})'
        processed=f'SUM(Variantenbilanz!E{start}:E{end})'
        investment=f'{calc(13)}*Szenarien!$B${r}'
        net=f'MAX(0,{investment}-MIN({inp("grant")},{investment}))'
        power=f'{processed}*{inp("energy_intensity")}'
        cstart=5+i*51; cend=cstart+50
        for x in range(cstart,cend+1):
            t=x-cstart; setval(cf,f'A{x}',name if t==0 else ''); setval(cf,f'B{x}',t,INT)
            if t==0: forms={'C':'0','D':'0','E':net,'F':'0','G':f'-E{x}','H':'1','I':f'G{x}','J':f'E{x}','K':'0'}
            else:
                active=f'AND(B{x}<={inp("horizon")},Szenarien!$H${r}="OK")'
                forms={
                 'C':f'IFERROR(IF({active},{used}*({inp("water_price")}*Szenarien!$C${r}*(1+{inp("water_growth")})^(B{x}-1)+{inp("sewage_price")}*{inp("sewage_credit")}*(1+{inp("sewage_growth")})^(B{x}-1)),0),0)',
                 'D':f'IFERROR(IF({active},(({inp("maintenance")}+{inp("other_cost")})*(1+{inp("cost_growth")})^(B{x}-1)+{power}*{inp("energy_price")}*(1+{inp("energy_growth")})^(B{x}-1))*Szenarien!$F${r},0),0)',
                 'E':f'IFERROR(IF(AND({active},B{x}={inp("replacement_year")}),{inp("replacement_cost")}*Szenarien!$B${r}*(1+{inp("cost_growth")})^(B{x}-1),0),0)',
                 'F':f'IFERROR(IF(AND({active},B{x}={inp("horizon")}),{inp("residual")}*Szenarien!$B${r},0),0)',
                 'G':f'C{x}-D{x}-E{x}+F{x}',
                 'H':f'IFERROR(IF({active},1/(1+{inp("discount")}+Szenarien!$G${r})^B{x},0),0)',
                 'I':f'G{x}*H{x}','J':f'(D{x}+E{x}-F{x})*H{x}',
                 'K':f'IF({active},{used}*H{x},0)'}
            for c,f in forms.items(): setval(cf,f'{c}{x}','='+f,NUM)
        usedref=f'I{r}'; quotaref=f'J{r}'; npv=f'L{r}'
        out={
        'I':used,'J':f'IFERROR(I{r}/({demand}+{inp("other_water")}),0)',
        'K':f'IFERROR(I{r}*({inp("water_price")}*C{r}+{inp("sewage_price")}*{inp("sewage_credit")})-({inp("maintenance")}+{inp("other_cost")}+{power}*{inp("energy_price")} )*F{r},0)',
        'L':f'SUM(VariantenCF!I{cstart}:I{cend})',
        'M':f'IF({net}=0,"Keine Nettoinvestition",IF(K{r}<=0,"Nicht erreichbar",{net}/K{r}))',
        'N':label(f'IF({net}=0,IF(L{r}>=0,1,3),IF(L{r}/{net}>=Schwellenwerte!B5,1,IF(L{r}/{net}>=Schwellenwerte!B6,2,3)))'),
        'O':label(f'IF(I{r}<=0,3,IF(J{r}>=Schwellenwerte!B9,1,IF(J{r}>=Schwellenwerte!B10,2,3)))'),
        'P':f'IF(OR(N{r}="Rot",O{r}="Rot",{calc(27)}=3),"Rot",IF(OR(N{r}="Gelb",O{r}="Gelb",{calc(27)}=2),"Gelb","Grün"))'}
        for c,f in out.items(): setval(ws,f'{c}{r}',f'=IF(H{r}="OK",{f},"Nicht bewertbar")',PCT if c=='J' else NUM)
        ws.row_dimensions[r].height=36
    for row,text in [(5,'Faktor 1,00 = tatsächlicher Projektansatz. Basisszenario startet mit 1,00; Projekt-Dashboard bleibt unverändert.'),(12,'Einzelfaktor-Sensitivitäten: alle anderen Faktoren = 1. Die Beschriftungen beschreiben die Startwerte; Werte frei editierbar.'),(27,'Keine Wahrscheinlichkeiten: Konservativ / Optimistisch sind Modellvarianten, keine Prognose oder statistische Unsicherheit.'),(28,'Andere Trinkwasser-Baseline außerhalb Betriebsbedarf bleibt konstant; Förderung absolut, Ersatzkosten/Restwert skaliert mit Investition.')]:
        ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=16); setval(ws,f'A{row}',text); ws.row_dimensions[row].height=30
    ws.column_dimensions['A'].width=29
    traffic(ws,'N8:P25')
    chart(ws,'bar','Kapitalwert der drei Szenarien',Reference(ws,min_col=12,min_row=7,max_row=10),Reference(ws,min_col=1,min_row=8,max_row=10),'A31','€',(17,9))
    chart(ws,'bar','Sensitivität des Kapitalwerts',Reference(ws,min_col=12,min_row=13,max_row=25),Reference(ws,min_col=1,min_row=14,max_row=25),'H31','€',(23,11))
    # Separate kleine Diagrammtabelle verhindert Titelübernahme aus leerer Zeile 13.
    ws['L13']='Kapitalwert €'; ws.print_area='A1:P54'; vws.print_area='A1:N25'; cf.print_area='A1:K25'

def make_dashboard(wb):
    ws=wb['Dashboard']; header(ws,'GRAUWASSER · Bewertungsdashboard','Wissenschaftlicher Excel-Prototyp · makrofrei · Eingaben blau · Ampelschwellen vorläufig · keine Freigabe',12)
    for c in range(1,13): ws.column_dimensions[col(c)].width=14
    ws.freeze_panes='A7'
    for row,key,labeltext in [(4,'name','Projekt'),(5,'building','Gebäudetyp')]:
        ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=2); setval(ws,f'A{row}',labeltext)
        ws.merge_cells(start_row=row,start_column=3,end_row=row,end_column=8); setval(ws,f'C{row}','='+inp(key)); ws.row_dimensions[row].height=24
    ws.merge_cells('I4:L4'); setval(ws,'I4',f'="Zeitraum: "&{inp("horizon")}&" Jahre | Bilanz "&{inp("year")}')
    ws.merge_cells('I5:L5'); setval(ws,'I5','='+calc(41)); ws.row_dimensions[5].height=35
    cards=[(7,1,'Grauwasseranfall',6,'#,##0.0" m³/a"'),(7,4,'Betriebswasserbedarf',9,'#,##0.0" m³/a"'),(7,7,'Nutzbar / Trinkwasser gespart',10,'#,##0.0" m³/a"'),(7,10,'Trinkwassereinsparquote',22,PCT),
           (11,1,'Gesamtinvestition',13,EUR),(11,4,'Betriebskosten Jahr 1',16,'#,##0.0" €/a"'),(11,7,'Vermiedene Kosten Jahr 1',17,'#,##0.0" €/a"'),(11,10,'Netto-Cashflow Jahr 1',18,'#,##0.0" €/a"'),
           (15,1,'Statische Amortisation',19,'0.0" Jahre"'),(15,4,'Dynamische Amortisation',21,'0.0" Jahre"'),(15,7,'Kapitalwert',20,EUR),(15,10,'Deckungsgrad',23,PCT)]
    for r,c,title,k,fmt in cards:
        ws.merge_cells(start_row=r,start_column=c,end_row=r,end_column=c+2)
        labelcell=setval(ws,f'{col(c)}{r}',title); labelcell.fill=PatternFill('solid',fgColor=NAVY); labelcell.font=Font(name='Aptos',size=10,color=WHITE,bold=True); ws.row_dimensions[r].height=28
        ws.merge_cells(start_row=r+1,start_column=c,end_row=r+2,end_column=c+2)
        cell=setval(ws,f'{col(c)}{r+1}',f'=IF({calc(5)}>0,"Eingaben prüfen",{calc(k)})',fmt)
        cell.fill=PatternFill('solid',fgColor=PALE); cell.font=Font(name='Aptos',size=18,bold=True,color=NAVY)
        cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        ws.row_dimensions[r+1].height=26; ws.row_dimensions[r+2].height=22
    for c,name,k in [(1,'Ökonomie',31),(5,'Ressourcen',32),(9,'Gesamtbewertung',33)]:
        ws.merge_cells(start_row=19,start_column=c,end_row=19,end_column=c+3); setval(ws,f'{col(c)}19',name)
        ws.merge_cells(start_row=20,start_column=c,end_row=20,end_column=c+3)
        cell=setval(ws,f'{col(c)}20','='+calc(k)); cell.font=Font(name='Aptos',size=20,bold=True,color=NAVY); cell.alignment=Alignment(horizontal='center',vertical='center'); traffic(ws,f'{col(c)}20:{col(c+3)}20')
    ws.row_dimensions[20].height=36
    reasons={
    21:f'=IF({calc(5)}>0,"Bewertung gesperrt: "&{calc(5)}&" Pflichtfehler; Blatt Prüfungen öffnen.","Ökonomie: Kapitalwert "&TEXT({calc(20)},"#,##0.00")&" €; Kriterium NPV/Nettoinvestition, Grün ab "&TEXT(Schwellenwerte!B5,"0.0%")&", Gelb ab "&TEXT(Schwellenwerte!B6,"0.0%")&". Ohne Nettoinvestition entscheidet Vorzeichen des NPV.")',
    22:f'=IF({calc(5)}>0,"Ressourcenbewertung gesperrt.","Ressourcen: Einsparquote "&TEXT({calc(22)},"0.0%")&"; Grün ab "&TEXT(Schwellenwerte!B9,"0.0%")&", Gelb ab "&TEXT(Schwellenwerte!B10,"0.0%")&". Ohne Einsparung Rot. Keine vollständige Umweltbilanz.")',
    23:f'="Gesamt = schlechteste Kategorie einschließlich Nachweisstatus. "&{calc(34)}&". Statische Amortisation nur innerhalb Horizont interpretieren; dynamische Stabilität: "&{calc(37)}&"."',
    24:'='+calc(40)}
    for r,f in reasons.items(): ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=12); setval(ws,f'A{r}',f); ws.row_dimensions[r].height=38 if r<24 else 48
    # Diagrammdaten sind sichtbar weiter unten; keine versteckten Konstanten oder Dateien.
    table(ws,46,['Wasserbilanz','m³/a'])
    for r,name,k in [(47,'Rohgrauwasser',6),(48,'Aufbereitet',8),(49,'Wiederverwendung',10),(50,'Betriebsbedarf',9),(51,'Nachspeisung',11)]:
        setval(ws,f'A{r}',name); setval(ws,f'B{r}',f'=IF({calc(5)}=0,{calc(k)},0)',NUM)
    table(ws,54,['Jahr','Cashflow €','Kumuliert diskontiert €'])
    for r in range(55,106):
        source=r-48; setval(ws,f'A{r}',f'=IFERROR(MIN(Zahlungen!A{source},{inp("horizon")}),0)',INT); setval(ws,f'B{r}',f'=IF(A{r}<={inp("horizon")},Zahlungen!G{source},NA())',NUM); setval(ws,f'C{r}',f'=IF(A{r}<={inp("horizon")},Zahlungen!J{source},NA())',NUM)
    # XY-Diagramme: Jahre oberhalb Horizont auf den Endpunkt klemmen.
    # Deckungsgleiche Endpunkte verändern Kurve/Achsen nicht; keine #N/A-Hilfszellen.
    for r in range(55,106):
        source=r-48
        ws[f'B{r}']=f'=IF({calc(5)}=0,INDEX(Zahlungen!$G$7:$G$57,A{r}+1),0)'
        ws[f'C{r}']=f'=IF({calc(5)}=0,INDEX(Zahlungen!$J$7:$J$57,A{r}+1),0)'
        ws[f'B{r}'].number_format=NUM; ws[f'C{r}'].number_format=NUM
    chart(ws,'bar','Jährliche Wasserbilanz',Reference(ws,min_col=2,min_row=46,max_row=51),Reference(ws,min_col=1,min_row=47,max_row=51),'A27','m³/a',(15,9))
    chart(ws,'scatter','Jährliche Zahlungsströme',Reference(ws,min_col=2,min_row=54,max_row=105),Reference(ws,min_col=1,min_row=55,max_row=105),'G27','€',(15,9))
    chart(ws,'scatter','Kumulierter diskontierter Cashflow',Reference(ws,min_col=3,min_row=54,max_row=105),Reference(ws,min_col=1,min_row=55,max_row=105),'A44','€',(30,8))
    # Diagrammtabellen abseits der Druckansicht platzieren.
    for row in list(range(46,52))+list(range(54,106)):
        for column in range(1,4):
            target=ws.cell(row,15+column); source=ws.cell(row,column)
            target.value=source.value; target.number_format=source.number_format
            source.value=None
    # Referenzen der bereits erzeugten Diagramme auf neue Tabellen setzen.
    for ch in ws._charts:
        for ser in ch.ser:
            if ser.val and ser.val.numRef: ser.val.numRef.f=ser.val.numRef.f.replace('$B$', '$Q$').replace('$C$', '$R$')
            if ser.yVal and ser.yVal.numRef: ser.yVal.numRef.f=ser.yVal.numRef.f.replace('$B$', '$Q$').replace('$C$', '$R$')
            if ser.xVal and ser.xVal.numRef: ser.xVal.numRef.f=ser.xVal.numRef.f.replace('$A$', '$P$')
            if ser.tx and ser.tx.strRef: ser.tx.strRef.f=ser.tx.strRef.f.replace('B54','Q54').replace('C54','R54').replace('B46','Q46').replace('$B$', '$Q$').replace('$C$', '$R$')
            if ser.cat and ser.cat.numRef: ser.cat.numRef.f=ser.cat.numRef.f.replace('$A$', '$P$')
        ch.display_blanks='gap'
    # Verschobene Diagrammtabellen: Jahresbezug jetzt in Spalte P.
    for r in range(55,106):
        ws[f'Q{r}']=str(ws[f'Q{r}'].value).replace(f'A{r}',f'P{r}')
        ws[f'R{r}']=str(ws[f'R{r}'].value).replace(f'A{r}',f'P{r}')
    ws.column_dimensions['P'].width=24; ws.column_dimensions['Q'].width=20; ws.column_dimensions['R'].width=25
    ws.merge_cells('A62:L62'); setval(ws,'A62','Anleitung → Eingaben → Profile → Kalender → Nachweise → Dashboard → Szenarien. Tabellen rechts liefern Diagrammdaten.'); ws.row_dimensions[62].height=30
    ws.print_area='A1:L63'; ws.page_setup.paperSize=ws.PAPERSIZE_A3

def make_kpis(wb):
    ws=wb['KPI_Katalog']; header(ws,'KPI-Katalog','Definitionen, Interpretation, Herkunft und Grenzen; nur ein Kriterium je Bewertungskategorie.',8)
    table(ws,4,['KPI','Definition','Formel','Einheit','Interpretation','Datenquelle','Grenzen','Ampel'])
    data=[
    ('Trinkwassereinsparung','Ersetzte Trinkwassermenge','Σ(Altverbrauch + Direktverbrauch)','m³/a','Absolute lokale Wassersubstitution','Profile, Kalender, Hersteller-/Speicherdaten','Keine LCA; Tagesmodell vereinfacht Zeitverlauf','Nullmenge → Ressourcen Rot; sonst informativ'),
    ('Einsparquote','Substitution bezogen auf Baseline','Nutzung / (Betriebsbedarf + sonstiger Trinkwasserbedarf)','%','Relative Verringerung des Trinkwasserbezugs','Zähler/Messung/Planung, keine Beispielwerte als Evidenz','Nenner beeinflusst Bewertung; Baseline konsistent halten','Einziges Ressourcen-Ampelkriterium'),
    ('Deckungsgrad','Deckung des Betriebsbedarfs','Nutzung / Betriebsbedarf','%','Anteil ohne Trinkwassernachspeisung','Nutzungsprofile, Tagesbilanz','Ohne Bedarf nicht definiert; intern 0 mit Hinweis','Informativ; mit Einsparquote korreliert'),
    ('GW-Nutzungsgrad','Anteil genutzter Rohmenge','Nutzung / Rohgrauwasser','%','Enthält Aufbereitungs- und Speicherverluste','Profile, Anlagenparameter','Hoher Anteil allein kein Eignungsnachweis; Nullnenner','Informativ; keine Doppelbewertung'),
    ('Gesamtinvestition','Brutto-Kapitaleinsatz','Planung + Anlage + Montage + Sonstiges','€','Kostenumfang, einheitlicher Preisstand','Angebote/Kostenschätzung; Beispiele nicht belegt','Ohne Kostenunsicherheit oder Finanzierung','Informativ'),
    ('Nettoinvestition','Eigenmittel im Jahr 0','Brutto − bestätigte Förderung','€','Förderung entlastet Kapitaleinsatz','Förderbescheid und Angebote','Förderung als sichere Jahr-0-Zahlung angenommen','Nenner relativer NPV, keine eigene Bewertung'),
    ('Nettoeinsparung','Wiederkehrender Netto-Cashflow Jahr 1','vermiedene variable Kosten − Betriebskosten','€/a','Erstes Betriebsjahr vor Ersatz und Zins','Tarife, Wartung, Energie, Mengen','Kein Äquivalent zum Kapitalwert','Informativ'),
    ('Statische Amortisation','Zeit bei konstantem Jahr-1-CF','Nettoinvestition / CF1, nur CF1 > 0','Jahre','Einfache Rückgewinnung; Horizont beachten','Kosten- und Nutzendaten','Kein Zins, keine Ersatzinvestition; außerhalb Horizont ausweisen','Informativ'),
    ('Kapitalwert','Diskontierter Gesamtnutzen','−I_net + Σ CF_t / (1+i)^t','€','≥0: wirtschaftlich unter Annahmen; relativ I_net klassifiziert','Nominale Kosten, nominaler Zins, Restwert','Zukunft unsicher; kein Risiko-/Steuermodell','Einziges Ökonomie-Ampelkriterium'),
    ('Dynamische Amortisation','Erste diskontierte Rückgewinnung','Erster kumulierter Barwert ≥0, linear interpoliert','Jahre','Zeit bis erste Deckung des eingesetzten Kapitals','Zahlungsreihe','Interpolation trotz Jahresendzahlungen; spätere Unterdeckung möglich','Informativ, Stabilitätswarnung ergänzend'),
    ('Spezifische Lebenszykluskosten','Diskontierte Kosten pro ersetztem m³','PV(I+OPEX+Ersatz−Restwert) / PV(Einsparmenge)','€/m³','Kosten der lokalen Wasserbereitstellung','Zahlungs- und Wasserbilanz','Kein Gebührenabzug im Zähler; PV-Volumen Konvention, keine Normformel','Informativ, kein zweites Ökonomiekriterium'),
    ('Strombedarf','Aufbereitungselektrizität','Verarbeitete Rohmenge × kWh/m³','kWh/a','Energieaufwand einschließlich ungenutzter Produktion','Hersteller-/Messdaten','Keine Pumpdetailrechnung, CO₂- oder Ökobilanz','Informativ')]
    for r,row in enumerate(data,5):
        for j,v in enumerate(row,1): setval(ws,f'{col(j)}{r}',v)
        ws.row_dimensions[r].height=84
    for c,w in {'A':26,'B':37,'C':53,'D':11,'E':43,'F':40,'G':50,'H':39}.items(): ws.column_dimensions[c].width=w
    ws.print_area='A1:H16'; ws.auto_filter.ref='A4:H16'

def make_user_test(wb):
    ws=wb['Nutzertest']; header(ws,'Anonymisierter Probandentest','Nur IDs erfassen, keine Namen. Nutzertest bewertet Verständlichkeit und Bedienung, keine fachliche Rechenvalidierung.',18)
    table(ws,5,['Anonyme ID','Rolle','Erfahrung Excel 1–5','Zeit Minuten','Fehlerzahl','Verständlichkeit','Bedienbarkeit','Übersichtlichkeit','Rechenweg verständlich','Ampel verständlich','Fehlersicherheit','Praxistauglichkeit','A1 Projekt erfasst','A2 Szenario erstellt','A3 Ampel erklärt','A4 Fehler erkannt','Beobachtete Fehler','Kommentar'])
    for r in range(6,56):
        for c in range(1,19): setval(ws,f'{col(c)}{r}',None,NUM if c==4 else INT if c in range(3,13) else None,True)
        dv(ws,f'B{r}',choices=['TGA-Planung','Studium','Betrieb','Sonstige'])
        for c in [3]+list(range(6,13)): dv(ws,f'{col(c)}{r}',1,5,True)
        dv(ws,f'D{r}',0,1440); dv(ws,f'E{r}',0,10000,True)
        for c in range(13,17): dv(ws,f'{col(c)}{r}',choices=['Erfüllt','Teilweise','Nicht erfüllt'])
    tab=Table(displayName='AnonymeTests',ref='A5:R55'); tab.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True); ws.add_table(tab)
    ws.column_dimensions['A'].width=20; ws.column_dimensions['Q'].width=44; ws.column_dimensions['R'].width=44
    ws.print_area='A1:R15'
    out=wb['Testauswertung']; header(out,'Nutzertest · Auswertung','Likert 1 = stimme gar nicht zu, 5 = stimme voll zu. Fehlende Antworten bleiben leer; keine imputierten Werte.',12)
    table(out,5,['Dimension','n gültig','Mittelwert','Streuung (Stichprobe)','Antwort 1','Antwort 2','Antwort 3','Antwort 4','Antwort 5','Fehlend bei ID','Ungültig bei ID','Interpretation'])
    questions=['Verständlichkeit','Bedienbarkeit','Übersichtlichkeit','Rechenweg verständlich','Ampel verständlich','Fehlersicherheit','Praxistauglichkeit']
    for r,(question,column) in enumerate(zip(questions,range(6,13)),6):
        ref=f'Nutzertest!{col(column)}$6:{col(column)}$55'; ids='Nutzertest!$A$6:$A$55'
        setval(out,f'A{r}',question)
        for c,rating in enumerate(range(1,6),5): setval(out,f'{col(c)}{r}',f'=SUMPRODUCT(({ids}<>"")*ISNUMBER({ref})*({ref}={rating}))',INT)
        setval(out,f'B{r}',f'=SUM(E{r}:I{r})',INT)
        setval(out,f'C{r}',f'=IF(B{r}=0,"",SUMPRODUCT(E{r}:I{r},$E$4:$I$4)/B{r})',NUM)
        setval(out,f'D{r}',f'=IF(B{r}<2,"",SQRT(SUMPRODUCT(E{r}:I{r},($E$4:$I$4-C{r})^2)/(B{r}-1)))',NUM)
        setval(out,f'J{r}',f'=COUNTIFS({ids},"<>",{ref},"")',INT)
        setval(out,f'K{r}',f'=COUNTIF({ids},"<>")-B{r}-J{r}',INT)
        setval(out,f'L{r}','Deskriptiv; keine fachliche Formelvalidierung')
    for c,rating in enumerate(range(1,6),5): setval(out,f'{col(c)}4',rating,INT)
    items={15:('Anzahl IDs','=COUNTIF(Nutzertest!A6:A55,"<>")'),16:('IDs mehrfach (Zeilen)','=SUMPRODUCT((Nutzertest!A6:A55<>"")*(COUNTIF(Nutzertest!A6:A55,Nutzertest!A6:A55)>1))'),17:('Gültige Bearbeitungszeiten n','=COUNTIFS(Nutzertest!A6:A55,"<>",Nutzertest!D6:D55,">=0",Nutzertest!D6:D55,"<=1440",Nutzertest!D6:D55,"<>")'),18:('Mittlere Zeit Minuten','=IF(B17=0,"",AVERAGEIFS(Nutzertest!D6:D55,Nutzertest!A6:A55,"<>",Nutzertest!D6:D55,">=0",Nutzertest!D6:D55,"<=1440"))'),19:('Mittlere Fehlerzahl','=IFERROR(AVERAGEIFS(Nutzertest!E6:E55,Nutzertest!A6:A55,"<>",Nutzertest!E6:E55,">=0",Nutzertest!E6:E55,"<=10000"),"")'),20:('Nicht zugeordnete Datenzeilen','=SUMPRODUCT((Nutzertest!A6:A55="")*((Nutzertest!F6:F55<>"")+(Nutzertest!D6:D55<>"")>0))')}
    ids='Nutzertest!$A$6:$A$55'; times='Nutzertest!$D$6:$D$55'; errors='Nutzertest!$E$6:$E$55'
    valid_time=f'({ids}<>"")*ISNUMBER({times})*({times}>=0)*({times}<=1440)'
    valid_errors=f'({ids}<>"")*ISNUMBER({errors})*({errors}>=0)*({errors}<=10000)*(MOD({errors},1)=0)'
    items[17]=('Gültige Bearbeitungszeiten n',f'=SUMPRODUCT(IFERROR({valid_time},0))')
    items[18]=('Mittlere Zeit Minuten',f'=IF(B17=0,"",SUMPRODUCT(IFERROR({valid_time}*{times},0))/B17)')
    items[19]=('Mittlere Fehlerzahl',f'=IF(B21=0,"",SUMPRODUCT(IFERROR({valid_errors}*{errors},0))/B21)')
    nonempty='+'.join(f'(Nutzertest!{col(c)}6:{col(c)}55<>"")' for c in range(2,19))
    items[20]=('Nicht zugeordnete Datenzeilen',f'=SUMPRODUCT(({ids}="")*(({nonempty})>0))')
    items[21]=('Gültige Fehlerzahlen n',f'=SUMPRODUCT(IFERROR({valid_errors},0))')
    items[22]=('Ungültige Zeit-/Fehlerfelder',f'=SUMPRODUCT(({ids}<>"")*({times}<>""))-B17+SUMPRODUCT(({ids}<>"")*({errors}<>""))-B21')
    for r,(name,f) in items.items(): setval(out,f'A{r}',name); setval(out,f'B{r}',f,NUM)
    table(out,23,['Aufgabe','Erfüllt','Teilweise','Nicht erfüllt'])
    for r,(name,column) in enumerate([('Projekt erfassen',13),('Szenario ändern',14),('Ampel begründen',15),('Fehler erkennen',16)],24):
        setval(out,f'A{r}',name)
        for j,status in enumerate(['Erfüllt','Teilweise','Nicht erfüllt'],2): setval(out,f'{col(j)}{r}',f'=COUNTIFS(Nutzertest!$A$6:$A$55,"<>",Nutzertest!{col(column)}$6:{col(column)}$55,"{status}")',INT)
    out.column_dimensions['A'].width=35
    chart(out,'bar','Mittlere Nutzerbewertungen',Reference(out,min_col=3,min_row=5,max_row=12),Reference(out,min_col=1,min_row=6,max_row=12),'A30','Skala 1–5',(15,9))
    chart(out,'bar','Antwortverteilungen',Reference(out,min_col=5,max_col=9,min_row=5,max_row=12),Reference(out,min_col=1,min_row=6,max_row=12),'G30','Anzahl',(17,9))
    for ch in out._charts:
        if ch.title and ch==out._charts[0]: ch.y_axis.scaling.min=1; ch.y_axis.scaling.max=5
    out.print_area='A1:L49'

def make_guide(wb):
    ws=wb['Anleitung']; header(ws,'Anleitung & wissenschaftliche Einordnung','Version 1.0 · 14.10.2026 · Bearbeitbare Annahmen, keine normative Freigabe oder Literaturvalidierung der Schwellen.',6)
    paragraphs=[
    ('1  Start','Nur blaue Felder bearbeiten. Eingaben, Profile und Kalender abgleichen. Das Dashboard rechnet automatisch; Excel-Berechnungsmodus Automatisch verwenden.'),
    ('2  Datenbasis','Auslieferung enthält ausschließlich fiktive Beispieldaten. Eine bestehende Excel-Datei wurde nicht bereitgestellt. Baseline, Angebote, Tarife und Herstellernachweise ergänzen.'),
    ('3  Kalender','Jede aktive Gruppe erhält gleiche Tagesmenge: Profil-Jahresmenge / aktive Tage. Profiltage können kleiner als Gruppentage sein. Ferien, Wochenenden und Hallenfremdnutzung manuell abbilden.'),
    ('4  Speicher','Tank startet leer. Nur frischer Tagesrest wird einen Tag übertragen; ungenutzte Altmenge verworfen. Dies ist eine Modellannahme, keine hygienisch zulässige Speicherzeit.'),
    ('5  Kosten','Einheitliche Währung und Preisbasis. Nur variable Wasser-/Abwassergebühren vermeiden. Abwasseranteil standardmäßig 0; Projektbeleg erforderlich. Nominalzins und nominale Steigerungen konsistent wählen.'),
    ('6  Ampel','Ökonomie: relativer Kapitalwert; Ressourcen: Einsparquote. Schlechteste Kategorie zählt. Nachweise ungeklärt: mindestens Gelb / Prüfbedarf; nicht erfüllt: Rot. Keine Kompensation.'),
    ('7  Amortisation','Statisch: Jahr-1-CF ohne Ersatz/Zins. Werte oberhalb Horizont sind keine im Projektzeitraum nachgewiesene Amortisation. Dynamisch: erste diskontierte Rückgewinnung; spätere Unterdeckung separat angezeigt.'),
    ('8  Szenarien','Drei Szenarien und zwölf Einzelfaktorvarianten sind konfigurierbar; erneute Tagesrechnung je Variante. Projekt-Dashboard bleibt auf Originaleingaben. Änderungen > Grenzbereich sperren nur betroffene Variante.'),
    ('9  Qualität','Prüfungen zeigt alle Pflichtfehler. Gültige Nullmengen sind zulässig und ergeben keinen Nutzen. Ungültige/fehlende Eingaben sperren das Dashboard; Hilfswerte sind dann nicht interpretierbar.'),
    ('10  Nutzerstudie','Nutzertest: 50 anonymisierte Personen, Likert 1–5, Aufgaben, Zeiten, Fehler. Auswertung nur mit ID; ungültige Ratings ausgeschlossen und gezählt. Duplikate separat prüfen.'),
    ('11  Grenzen','Keine stündliche Hydraulik, Wasserqualitätsberechnung, Ausfallsimulation, Mehrfach-Ersatzzyklen, Finanzierung/Steuern oder vollständige Ökobilanz. Kosten-/Zeitannahmen projektspezifisch validieren.'),
    ('12  Regelwerke','DIN EN 16941-2 (aufbereitetes Grauwasser) und DIN EN 1717 (Schutz des Trinkwassers) als Recherchehinweise. Volltexte/Projektanwendbarkeit nicht geprüft; keine Normgrenzwerte übernommen.'),
    ('13  Evidenz','Alle Mengen, Preise, Anlagenparameter und Schwellen im Beispiel sind unbestätigte Modellannahmen. Status Nachgewiesen nur mit nachprüfbaren Projektbelegen setzen.'),
    ('14  Reproduzierbarkeit','Python-Generator und Tests liegen bei. Excel-Formeln verwenden Standardfunktionen ohne Makros, externe Links oder Datenverbindungen. LibreOffice-Enginevalidierung ersetzt keinen abschließenden Windows-Excel-Test.'),
    ('15  Begleitdateien','Methodik_Dokumentation.md enthält Gleichungen und KPI-Abgrenzungen; Probandentest.md enthält Aufgaben/Skala; Validierungsbericht.md dokumentiert tatsächlich ausgeführte Tests.')]
    for r,(title,text) in enumerate(paragraphs,5):
        setval(ws,f'A{r}',title); ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=6); setval(ws,f'B{r}',text); ws.row_dimensions[r].height=64
    ws.column_dimensions['A'].width=27
    for c in 'BCDEF': ws.column_dimensions[c].width=23
    ws.print_area='A1:F19'

def build(path):
    p=Project(); wb=Workbook(); wb.remove(wb.active)
    for name in NAMES: wb.create_sheet(name)
    wb.properties.creator='Grauwasser-Bewertungstool'; wb.properties.title='Bewertungssystem für Grauwasseranlagen'; wb.properties.subject='Wissenschaftlicher, makrofreier Excel-Prototyp'; wb.properties.description='Tagesbilanz, Kapitalwert, konservative Ampel, Szenarien, Probandenauswertung. Beispieldaten und vorläufige Schwellen.'
    wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True,forceFullCalc=True,calcMode='auto')
    make_inputs(wb,p); make_profiles(wb,p); make_calendar(wb,p); make_thresholds(wb,p)
    make_checks(wb,p); make_water(wb); make_payments(wb); make_calc(wb); make_scenarios(wb)
    make_dashboard(wb); make_kpis(wb); make_user_test(wb); make_guide(wb)
    for ws in wb:
        ws.sheet_view.zoomScale=85 if ws.title!='Dashboard' else 80
        # Schutz dient nur der Bedienung, nicht dem Schutz von Daten/Geheimnissen.
        ws.protection.sheet=True; ws.protection.selectLockedCells=False; ws.protection.selectUnlockedCells=False
        ws.protection.autoFilter=False; ws.protection.sort=False; ws.protection.objects=True
    wb.active=0
    path=Path(path); path.parent.mkdir(parents=True,exist_ok=True); wb.save(path)
    return path

if __name__=='__main__':
    parser=argparse.ArgumentParser(); parser.add_argument('--output',default='Grauwasser_Bewertungstool.xlsx'); args=parser.parse_args()
    print(build(args.output))
